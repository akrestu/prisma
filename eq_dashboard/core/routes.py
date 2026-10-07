"""Haul Routes workbook of one site: destinations (Tujuan) and routes per loader + destination (loading pit,
horizontal and vertical distance, valid from a date until a newer row replaces it). Hourly Production lines pick a
destination; pit and distances come from the route. No Streamlit or database here."""
from __future__ import annotations

import datetime as dt
import io
from dataclasses import dataclass, field

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from core.clean import _ids
from core.io import excel_date, frame, num, read_workbook, text
from core.validate import StructureError, file_stem

DATASET = "Haul Routes"
DEST_SHEET, ROUTE_SHEET = "Destinations", "Routes"
SITE_CELL = "B2"
HEADER_ROW = 4
GROUPS = ("OB", "CG")
DEST_COLS = ["name", "material_group", "active"]
ROUTE_COLS = ["loader", "destination", "pit", "dist_h", "dist_v", "valid_from"]
DEST_HEADS = {"Destination": "name", "Material": "material_group", "Active": "active"}
ROUTE_HEADS = {"Loader": "loader", "Destination": "destination", "PIT": "pit", "Horizontal (m)": "dist_h",
               "Vertical (m)": "dist_v", "Valid from": "valid_from"}


def file_name(site: str) -> str:
    """Haul_Routes_WBK-BAU.xlsx"""
    return f"{file_stem(DATASET)}_{site}.xlsx"


def _yes(v) -> bool:
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return True                                   # empty = active
    if isinstance(v, bool):
        return v
    return str(v).strip().lower() not in ("no", "n", "false", "0", "tidak", "inactive")


# ------------------------------------------------------------------ checking what was typed or uploaded
def clean_destinations(df: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """name, material_group (OB | CG), active → cleaned rows + problems (blocking)."""
    d = df.reindex(columns=DEST_COLS).copy()
    d["name"] = text(d["name"].astype("string"))
    d = d[d["name"].notna()].copy()
    d["material_group"] = d["material_group"].astype("string").str.strip().str.upper().str[:2]
    d["active"] = d["active"].map(_yes).astype(bool)
    problems = [f"Destination {n}: material must be OB or CG." for n in d.loc[~d["material_group"].isin(GROUPS), "name"]]
    dup = d[d["name"].str.upper().duplicated(keep=False)]
    if len(dup):
        problems.append(f"Destination listed twice: {', '.join(sorted(set(dup['name'])))}.")
    too_long = d[d["name"].str.len() > 120]
    if len(too_long):
        problems.append(f"Destination name longer than 120 characters: {', '.join(too_long['name'].str[:30])}…")
    return d.reset_index(drop=True), problems


def clean_routes(df: pd.DataFrame, destinations: pd.DataFrame, today: dt.date) -> tuple[pd.DataFrame, list[str]]:
    """loader, destination, pit, dist_h, dist_v, valid_from → cleaned rows + problems. An empty 'valid from' means
    today. The destination must be one of the site's destinations."""
    d = df.reindex(columns=ROUTE_COLS).copy()
    d["loader"] = _ids(d["loader"])
    d["destination"] = text(d["destination"].astype("string"))
    d["pit"] = text(d["pit"].astype("string"))
    d = d[d["loader"].notna() | d["destination"].notna()].copy()
    for c in ("dist_h", "dist_v"):
        d[c] = num(d[c])
    vf = excel_date(d["valid_from"]) if len(d) else pd.Series(dtype="datetime64[ns]")
    d["valid_from"] = [x.date() if pd.notna(x) else today for x in vf]
    known = {str(n).upper(): n for n in destinations["name"]} if len(destinations) else {}
    d["destination"] = d["destination"].map(lambda x: known.get(str(x).upper(), x) if pd.notna(x) else x)
    problems = []
    for i, r in enumerate(d.itertuples(), 1):
        where = (f"Route {i} ({r.loader if pd.notna(r.loader) else 'no loader'} → "
                 f"{r.destination if pd.notna(r.destination) else 'no destination'})")
        if pd.isna(r.loader):
            problems.append(f"{where}: loader is empty.")
        if pd.isna(r.destination):
            problems.append(f"{where}: destination is empty.")
        elif str(r.destination).upper() not in known:
            problems.append(f"{where}: {r.destination} is not in the destinations of this site; add it there first.")
        for c, lab in (("dist_h", "horizontal"), ("dist_v", "vertical")):
            v = getattr(r, c)
            if pd.notna(v) and not (0 <= v <= 50_000):
                problems.append(f"{where}: {lab} distance must be between 0 and 50,000 m.")
    dup = d[d.duplicated(["loader", "destination", "valid_from"], keep=False)]
    if len(dup):
        pairs = sorted({f"{a} → {b} ({c:%d %b %Y})" for a, b, c in dup[["loader", "destination", "valid_from"]]
                        .itertuples(index=False)})
        problems.append(f"Same route and date twice: {', '.join(pairs)}.")
    return d.reset_index(drop=True), problems


def routes_at(routes: pd.DataFrame, date: dt.date) -> pd.DataFrame:
    """The route in force on `date` for each loader + destination: the newest row valid on or before it."""
    if routes is None or routes.empty:
        return pd.DataFrame(columns=ROUTE_COLS)
    r = routes[pd.to_datetime(routes["valid_from"]).dt.date <= date]
    return (r.sort_values("valid_from").drop_duplicates(["loader", "destination"], keep="last")
            .reset_index(drop=True))


def current(routes: pd.DataFrame, date: dt.date) -> pd.DataFrame:
    """Routes in force on `date` plus those that start later (planned): what the setup grid shows and edits."""
    if routes is None or routes.empty:
        return pd.DataFrame(columns=ROUTE_COLS)
    later = routes[pd.to_datetime(routes["valid_from"]).dt.date > date]
    return (pd.concat([routes_at(routes, date), later], ignore_index=True)
            .sort_values(["loader", "destination", "valid_from"]).reset_index(drop=True))


def merge_routes(stored: pd.DataFrame, edited: pd.DataFrame, date: dt.date) -> pd.DataFrame:
    """Rows kept after a save: routes already replaced before `date` stay as history; the grid (routes in force on
    `date` and later) replaces the rest. A new distance with a later 'valid from' therefore keeps the old row for
    the shifts before it."""
    edited = edited.reindex(columns=ROUTE_COLS)
    if stored is None or stored.empty:
        return edited
    key = ["loader", "destination", "valid_from"]
    shown = set(map(tuple, current(stored, date)[key].astype(str).to_numpy()))
    history = stored[[tuple(x) not in shown for x in stored[key].astype(str).to_numpy()]]
    out = pd.concat([history.reindex(columns=ROUTE_COLS), edited], ignore_index=True)
    return out.drop_duplicates(key, keep="last").sort_values(key).reset_index(drop=True)


# ------------------------------------------------------------------ workbook
def build_template(site: str, destinations: pd.DataFrame, routes: pd.DataFrame, loaders: list[str]) -> bytes:
    from core import dataprod
    wb = Workbook()
    ds = wb.active
    ds.title = DEST_SHEET
    _head(ds, site, "Where haulers unload: disposals for OB, ROM / stockpiles for coal. Active 'No' hides it from "
                    "the shift form without deleting it.", list(DEST_HEADS))
    for i, r in enumerate(destinations.reindex(columns=DEST_COLS).itertuples(index=False), start=HEADER_ROW + 1):
        ds.cell(i, 1, r.name)
        ds.cell(i, 2, r.material_group)
        ds.cell(i, 3, "Yes" if _yes(r.active) else "No")
    _list(ds, "B", '"OB,CG"')
    _list(ds, "C", '"Yes,No"')
    rs = wb.create_sheet(ROUTE_SHEET)
    _head(rs, site, "One row per loader + destination. A new distance from a date: add a row with that 'Valid from' "
                    "(older shifts keep the old distance).", list(ROUTE_HEADS))
    for i, r in enumerate(routes.reindex(columns=ROUTE_COLS).itertuples(index=False), start=HEADER_ROW + 1):
        for j, v in enumerate(r, start=1):
            rs.cell(i, j, None if v is None or (not isinstance(v, str) and pd.isna(v)) else v)
        rs.cell(i, 6).number_format = "yyyy-mm-dd"
    lists = wb.create_sheet("Lists")
    for i, v in enumerate(loaders, start=1):
        lists[f"A{i}"] = v
    names = destinations["name"].tolist() if len(destinations) else []
    for i, v in enumerate(names, start=1):
        lists[f"B{i}"] = v
    lists.sheet_state = "hidden"
    if loaders:
        _list(rs, "A", f"=Lists!$A$1:$A${len(loaders)}")
    if names:
        _list(rs, "B", f"=Lists!$B$1:$B${len(names)}")
    dataprod._meta(wb, "haul_routes", {"dataset": DATASET, "site": site})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _head(ws, site: str, tip: str, heads: list[str]) -> None:
    from core import dataprod
    ws["A1"] = f"PRISMA · {DATASET} · {ws.title}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"], ws[SITE_CELL] = "Site", site
    ws["A2"].font = Font(bold=True)
    ws["D2"] = tip
    ws["D2"].font = Font(italic=True, color="55595F")
    for j, h in enumerate(heads, start=1):
        c = ws.cell(HEADER_ROW, j, h)
        c.fill, c.font = dataprod.HEAD_FILL, dataprod.HEAD_FONT
        ws.column_dimensions[c.column_letter].width = 26 if h in ("Destination", "PIT") else 14
    ws.freeze_panes = ws.cell(HEADER_ROW + 1, 1)


def _list(ws, col: str, formula: str) -> None:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, errorStyle="warning")
    dv.add(f"{col}{HEADER_ROW + 1}:{col}{HEADER_ROW + 500}")
    ws.add_data_validation(dv)


@dataclass
class RouteFile:
    site: str | None
    destinations: pd.DataFrame
    routes: pd.DataFrame
    problems: list[str] = field(default_factory=list)


def _sheet(raw: dict, name: str, heads: dict) -> tuple[str | None, pd.DataFrame]:
    x = raw[name]
    site = str(x.iloc[1, 1]).strip() if len(x) > 1 and x.shape[1] > 1 and pd.notna(x.iloc[1, 1]) else None
    if len(x) < HEADER_ROW:
        return site, pd.DataFrame(columns=list(heads.values()))
    g = frame(x, HEADER_ROW - 1)
    missing = [h for h in heads if h not in g.columns]
    if missing:
        raise StructureError([f"'{name}': row {HEADER_ROW} must hold the headers {', '.join(heads)} "
                              f"(missing {', '.join(missing)}). Download the current template."])
    return site, g.rename(columns=heads)[list(heads.values())].dropna(how="all")


def parse(data: bytes, today: dt.date) -> RouteFile:
    raw = read_workbook(data)
    if DEST_SHEET not in raw:
        raise StructureError([f"Sheet '{DEST_SHEET}' was not found: use the {DATASET} template."])
    site, dest = _sheet(raw, DEST_SHEET, DEST_HEADS)
    rt = _sheet(raw, ROUTE_SHEET, ROUTE_HEADS)[1] if ROUTE_SHEET in raw else pd.DataFrame(columns=ROUTE_COLS)
    dest, p1 = clean_destinations(dest)
    rt, p2 = clean_routes(rt, dest, today)
    return RouteFile(site, dest, rt, p1 + p2)
