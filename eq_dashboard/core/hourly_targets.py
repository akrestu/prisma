"""Productivity targets workbook of one site: targets per excavator model and per hauler model (internal and
client) and excavator overrides. Download pre-filled, edit in Excel, upload back; the grids in Setup → Productivity
targets do the same.
No Streamlit or database here."""
from __future__ import annotations

import io
from dataclasses import dataclass, field

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from core import prod_target as PT
from core.clean import _ids
from core.io import frame, read_workbook
from core.validate import StructureError, file_stem

DATASET = "Productivity targets"
MODELS, HAULERS, UNITS = "Model Targets", "Hauler Targets", "Unit Overrides"
MODEL_VALUES = {"ob": "OB BCM/h", "mud": "Mud BCM/h", "coal": "Coal t/h"}
HAULER_VALUES = {"ob": "OB BCM/h", "coal": "Coal t/h"}
UNIT_COLS = ["Excavator", "Model", "Material", "Basis", "Target per hour"]
SITE_CELL = "B2"
HEADER_ROW = 4          # Excel row of the model table header


def file_name(site: str) -> str:
    """Productivity_targets_WBK-BAU.xlsx"""
    return f"{file_stem(DATASET)}_{site}.xlsx"


def build_template(site: str, models: pd.DataFrame, overrides: pd.DataFrame, site_models: list[str],
                   haulers: pd.DataFrame | None = None, hauler_models: list[str] | None = None) -> bytes:
    """`models`: excavator targets (model, basis, ob, mud, coal); every excavator model of the site gets a row.
    `haulers`: hauler targets (model, basis, ob, coal), a row per hauler model of the site. `overrides`: unit_id,
    model, material_group, basis, target_per_hour."""
    from core import dataprod
    wb = Workbook()
    ws = wb.active
    ws.title = MODELS
    ws["A1"] = f"PRISMA · {DATASET} · per excavator model"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"], ws[SITE_CELL] = "Site", site
    ws["A2"].font = Font(bold=True)
    ws["D2"] = ("Internal counts first, client where internal is empty. Mud is used for mud and mud blending; "
                "empty Mud or Coal = the OB value. A model without any value has no target.")
    ws["D2"].font = Font(italic=True, color="55595F")
    _model_sheet(ws, models, MODEL_VALUES, site_models)
    wh = wb.create_sheet(HAULERS)
    wh["A1"] = f"PRISMA · {DATASET} · per hauler model"
    wh["A1"].font = Font(bold=True, size=14)
    wh["A2"], wh[SITE_CELL] = "Site", site
    wh["A2"].font = Font(bold=True)
    _model_sheet(wh, haulers, HAULER_VALUES, hauler_models or [])

    wu = wb.create_sheet(UNITS)
    for j, h in enumerate(UNIT_COLS, start=1):
        c = wu.cell(1, j, h)
        c.fill, c.font = dataprod.HEAD_FILL, dataprod.HEAD_FONT
        wu.column_dimensions[c.column_letter].width = 16
    for r in overrides.itertuples():
        wu.append([r.unit_id, r.model, r.material_group, r.basis, r.target_per_hour])
    for col, choices in (("C", "OB,CG"), ("D", ",".join(PT.BASES))):
        v = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
        v.add(f"{col}2:{col}500")
        wu.add_data_validation(v)
    wu.freeze_panes = "A2"
    dataprod._meta(wb, "hourly_targets", {"dataset": DATASET, "site": site})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _model_sheet(ws, table: pd.DataFrame | None, values: dict[str, str], site_models: list[str]) -> None:
    """Model table from HEADER_ROW: model, then each value per basis; a row for every model of the site."""
    from core import dataprod
    grid = PT.wide(table, values)
    names = sorted(set(grid["model"]) | {m for m in site_models if PT.match(m, grid["model"]) is None})
    grid = grid.set_index("model").reindex(names).reset_index()
    for j, h in enumerate(grid.columns, start=1):
        c = ws.cell(HEADER_ROW, j, "Model" if h == "model" else h)
        c.fill, c.font = dataprod.HEAD_FILL, dataprod.HEAD_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = 18 if h == "model" else 14
    for i, r in enumerate(grid.itertuples(index=False), start=HEADER_ROW + 1):
        for j, v in enumerate(r, start=1):
            ws.cell(i, j, None if pd.isna(v) else v)
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="100000", allow_blank=True)
    dv.error, dv.errorTitle = "Enter a number ≥ 0", "Target"
    dv.add(f"B{HEADER_ROW + 1}:{chr(ord('A') + len(grid.columns) - 1)}{HEADER_ROW + 300}")
    ws.add_data_validation(dv)
    ws.freeze_panes = ws.cell(HEADER_ROW + 1, 2)


@dataclass
class TargetFile:
    site: str
    models: pd.DataFrame                      # model, basis, ob, mud, coal
    overrides: pd.DataFrame                   # unit_id, model, material_group, basis, target_per_hour
    problems: list[str] = field(default_factory=list)
    haulers: pd.DataFrame | None = None       # model, basis, ob, coal (None: older file without the sheet)


def _read_models(x: pd.DataFrame, sheet: str, values: dict[str, str]) -> pd.DataFrame:
    if len(x) < HEADER_ROW:
        raise StructureError([f"'{sheet}': the table header belongs in row {HEADER_ROW}."])
    g = frame(x, HEADER_ROW - 1).rename(columns={"Model": "model"})
    want = [f"{lbl} · {b}" for b in PT.BASES for lbl in values.values()]
    missing = [c for c in ["model", *want] if c not in g.columns]
    if missing:
        raise StructureError([f"'{sheet}' (row {HEADER_ROW}) is missing columns: {', '.join(missing)}."])
    for c in want:
        g[c] = pd.to_numeric(g[c], errors="coerce")
    return PT.long(g[["model", *want]], values)


def parse_template(data: bytes) -> TargetFile:
    raw = read_workbook(data)
    if MODELS not in raw:
        raise StructureError([f"Sheet '{MODELS}' was not found: use the {DATASET} template of the site."])
    x = raw[MODELS]
    site = str(x.iloc[1, 1]).strip() if len(x) > 1 and x.shape[1] > 1 and pd.notna(x.iloc[1, 1]) else ""
    if not site:
        raise StructureError([f"'{MODELS}': the site (cell {SITE_CELL}) is empty."])
    models = _read_models(x, MODELS, MODEL_VALUES)
    haulers = _read_models(raw[HAULERS], HAULERS, HAULER_VALUES) if HAULERS in raw else None
    problems = []
    over = pd.DataFrame(columns=["unit_id", "model", "material_group", "basis", "target_per_hour"])
    if UNITS in raw and len(raw[UNITS]):
        u = frame(raw[UNITS], 0)
        miss = [c for c in UNIT_COLS if c not in u.columns]
        if miss:
            raise StructureError([f"'{UNITS}' (row 1) is missing columns: {', '.join(miss)}."])
        u = u.dropna(subset=["Excavator"])
        over = pd.DataFrame({"unit_id": _ids(u["Excavator"]), "model": u["Model"].astype("string").str.strip(),
                             "material_group": u["Material"].astype("string").str.strip().str.upper(),
                             "basis": u["Basis"].astype("string").str.strip().str.lower(),
                             "target_per_hour": pd.to_numeric(u["Target per hour"], errors="coerce"),
                             "_row": u["_row"]})
        for r in over.itertuples():
            if r.material_group not in ("OB", "CG"):
                problems.append(f"'{UNITS}' row {r._row}: Material must be OB or CG.")
            if r.basis not in PT.BASES:
                problems.append(f"'{UNITS}' row {r._row}: Basis must be {' or '.join(PT.BASES)}.")
            if pd.isna(r.target_per_hour) or r.target_per_hour <= 0:
                problems.append(f"'{UNITS}' row {r._row}: Target per hour must be a number above 0.")
        dup = over[over.duplicated(["unit_id", "material_group", "basis"], keep=False)]
        if len(dup):
            problems.append(f"'{UNITS}': more than one target for {', '.join(sorted(set(dup['unit_id'])))} "
                            "with the same material and basis.")
        over = over.drop(columns="_row")
    return TargetFile(site, models, over.reset_index(drop=True), problems, haulers)
