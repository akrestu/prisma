"""Load Factors workbook of one site: load per trip by material × hauler model (BCM for OB, ton for coal), the
numbers that turn Hourly Production trips into volume. It replaces the old Mst Hourly workbook, whose 'Link Muatan'
sheet is still read. No Streamlit or database here."""
from __future__ import annotations

import io
from dataclasses import dataclass

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.clean import material_group
from core.io import frame, read_workbook
from core.validate import StructureError, file_stem

DATASET = "Load Factors"
SHEET = DATASET
LEGACY_SHEET = "Link Muatan"          # Mst Hourly workbook
SITE_CELL = "B2"
HEADER_ROW = 4


def file_name(site: str) -> str:
    """Load_Factors_WBK-BAU.xlsx"""
    return f"{file_stem(DATASET)}_{site}.xlsx"


def wide(lf: pd.DataFrame, models: list[str]) -> pd.DataFrame:
    """Long load factors → one row per material, one column per hauler model (`models` first, then any others)."""
    extra = [mo for mo in sorted(lf["hauler_model"].unique()) if mo not in models] if len(lf) else []
    cols = list(models) + extra
    if lf.empty:
        return pd.DataFrame(columns=["material", *cols]).astype(dict.fromkeys(cols, float))
    g = lf.pivot_table(index="material", columns="hauler_model", values="muatan", aggfunc="first").reindex(columns=cols)
    return g.reset_index().rename_axis(columns=None)


def long(grid: pd.DataFrame) -> pd.DataFrame:
    """Matrix → material, material_group, hauler_model, muatan (empty and non-positive cells dropped)."""
    g = grid.dropna(subset=["material"])
    out = g.melt(id_vars=["material"], var_name="hauler_model", value_name="muatan")
    out["muatan"] = pd.to_numeric(out["muatan"], errors="coerce")
    out = out[out["muatan"] > 0].copy()
    out["material"] = out["material"].astype(str).str.strip()
    out["hauler_model"] = out["hauler_model"].astype(str).str.strip()
    out.insert(1, "material_group", material_group(out["material"]).to_numpy())
    return out.drop_duplicates(["material", "hauler_model"], keep="last").reset_index(drop=True)


def build_template(site: str, lf: pd.DataFrame, models: list[str]) -> bytes:
    """Pre-filled matrix for the site; the columns are the truck models of its unit population."""
    from core import dataprod
    grid = wide(lf, models)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws["A1"] = f"PRISMA · {DATASET} · load per trip (BCM for OB, ton for coal)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"], ws[SITE_CELL] = "Site", site
    ws["A2"].font = Font(bold=True)
    ws["D2"] = ("One row per material, one column per truck model. Empty cell = that model does not carry that "
                "material. Material names start with OB or CG.")
    ws["D2"].font = Font(italic=True, color="55595F")
    for j, h in enumerate(grid.columns, start=1):
        c = ws.cell(HEADER_ROW, j, "Material" if h == "material" else h)
        c.fill, c.font = dataprod.HEAD_FILL, dataprod.HEAD_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = 22 if h == "material" else 12
    for i, r in enumerate(grid.itertuples(index=False), start=HEADER_ROW + 1):
        for j, v in enumerate(r, start=1):
            ws.cell(i, j, None if pd.isna(v) else v)
    if len(grid.columns) > 1:
        dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1000", allow_blank=True)
        dv.error, dv.errorTitle = "Load per trip: a number above 0", "Load"
        dv.add(f"B{HEADER_ROW + 1}:{get_column_letter(len(grid.columns))}{HEADER_ROW + 200}")
        ws.add_data_validation(dv)
    ws.freeze_panes = ws.cell(HEADER_ROW + 1, 2)
    dataprod._meta(wb, "load_factors", {"dataset": DATASET, "site": site})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@dataclass
class LoadFile:
    site: str | None          # None for a Mst Hourly workbook (it names no site)
    load: pd.DataFrame        # material, material_group, hauler_model, muatan
    source: str               # "Load Factors" | "Mst Hourly (Link Muatan)"
    legacy_targets: int = 0   # excavator targets found in a Mst Hourly file (not imported here)


def parse(data: bytes) -> LoadFile:
    raw = read_workbook(data)
    if SHEET in raw:
        x = raw[SHEET]
        site = str(x.iloc[1, 1]).strip() if len(x) > 1 and x.shape[1] > 1 and pd.notna(x.iloc[1, 1]) else None
        if len(x) < HEADER_ROW:
            raise StructureError([f"'{SHEET}': the table header belongs in row {HEADER_ROW}."])
        g = frame(x, HEADER_ROW - 1).drop(columns="_row").rename(columns={"Material": "material"})
        if "material" not in g.columns:
            raise StructureError([f"'{SHEET}': cell A{HEADER_ROW} must be 'Material' with truck models to its right."])
        g = g[[c for c in g.columns if not str(c).startswith("_col")]]
        load = long(g)
        source = DATASET
        legacy = 0
    elif LEGACY_SHEET in raw:
        from core.hourly import parse_link_muatan
        load, targets = parse_link_muatan(data)
        site, source, legacy = None, "Mst Hourly (Link Muatan)", len(targets)
    else:
        raise StructureError([f"Sheet '{SHEET}' was not found (a Mst Hourly workbook with '{LEGACY_SHEET}' also "
                              "works)."])
    bad = sorted(load.loc[load["material_group"] == "Other", "material"].unique()) if len(load) else []
    if bad:
        raise StructureError([f"Material must start with OB or CG: {', '.join(bad)}."])
    if load.empty:
        raise StructureError(["No load per trip found in the file."])
    return LoadFile(site, load, source, legacy)
