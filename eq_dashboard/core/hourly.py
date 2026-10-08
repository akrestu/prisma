"""Hourly Production (flash data, no approval): shift slots, the current production hour, master data from the
'Link Muatan' sheet, the per-shift Excel template, and turning input rows into volumes. No Streamlit or database.

A shift sheet has one line per hauler (with its loader and both operators) and trips per hour (12 slots). Volume = trips × load
factor (material × hauler model): BCM for OB, ton for coal. Hourly figures are an estimate for monitoring; the
official monthly numbers stay those of the approved Production Data.
"""
from __future__ import annotations

import datetime as dt
import io
from dataclasses import dataclass, field

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from core.clean import _ids, material_group
from core.io import excel_date, num, read_workbook, text
from core.validate import HOUR_SLOTS, HOURLY_PRODUCTION, StructureError, file_stem

SHIFTS = ("DS", "NS")
SLOTS = {"DS": HOUR_SLOTS[:12], "NS": HOUR_SLOTS[12:]}      # DS 06-07 … 17-18, NS 18-19 … 05-06
R = [f"r{i}" for i in range(1, 13)]                          # storage columns for the 12 slots
DATASET = HOURLY_PRODUCTION
SHEET = HOURLY_PRODUCTION                                    # "Hourly Production"; files made before used "Hourly"
SHEET_NAMES = (SHEET, "Hourly")
HEADER_ROW = 7                                               # Excel row of the table header in the template
# the order follows the work: excavator and its operator, material, then each truck it loads and where it goes. Only
# what the data officer knows goes in the file; hauler model and load come from the unit population. PIT, disposal
# and the H / V distances change with the front and the dump, so they are typed per line (the disposal from the
# site's list, the distances by the engineering checker).
INPUT_COLS = ["Loader", "Operator", "Material", "Hauler ID", "Hauler operator", "PIT", "Disposal", "H distance (m)",
              "V distance (m)"]
INPUT_COLS_V3 = ["Loader", "Operator", "Material", "Hauler ID", "Hauler operator", "Destination", "H distance (m)",
                 "V distance (m)"]
INPUT_COLS_V2 = ["Loader", "Operator", "Material", "Hauler ID", "Hauler operator", "Destination"]
INPUT_COLS_V1 = ["Loader", "Operator", "Material", "Hauler ID", "Hauler operator", "PIT", "Disposal", "Distance (m)"]
MAX_TRIPS = 20                                               # trips of one hauler in one hour, all its lines together
TAIL_COLS = ["Remark code", "Remark"]                       # old templates only: one remark per line (ignored now)
REMARK_SHEET = "Remarks"
REMARK_HEADS = ["Hour", "Loader", "Hauler ID", "Remark code", "Remark"]
REMARKS = {  # common reasons on the site boards; any 3-digit Equipment Events reason code is also accepted
    "100": "Productivity achieved", "101": "Change shift", "301": "Standby / waiting", "302": "Rain",
    "303": "Slippery road", "304": "Blasting", "305": "Waiting hauler", "401": "Breakdown loader",
    "402": "Breakdown hauler", "501": "Scheduled maintenance", "502": "Digging method / front condition",
}


def clean_remarks(df: pd.DataFrame, shift: str, loaders: list[str] | None = None) -> tuple[pd.DataFrame, list[str]]:
    """Remarks per hour from the web table or the Remarks sheet → (slot, loader, hauler, code, remark), problems.

    Hour may be the slot label ('09-10') or 1..12. The code is the 3 digits in front ('302 - Rain' → '302'). Empty
    lines are dropped; a line needs an hour, a loader and a code or a text."""
    cols = ["slot", "loader", "hauler", "code", "remark"]
    if df is None or df.empty:
        return pd.DataFrame(columns=cols), []
    d = df.reindex(columns=["hour", "loader", "hauler", "code", "remark"]).copy()
    d = d[d.apply(lambda r: any(pd.notna(v) and str(v).strip() for v in r), axis=1)]
    if d.empty:
        return pd.DataFrame(columns=cols), []
    labels = {lab: k for k, lab in enumerate(SLOTS[shift], 1)}

    def to_slot(v):
        if v is None or pd.isna(v):
            return None
        h = str(v).strip()
        if h in labels:
            return labels[h]
        try:
            k = int(float(h))
        except ValueError:
            return None
        return k if 1 <= k <= 12 else None

    out = pd.DataFrame({
        "slot": d["hour"].map(to_slot),
        "loader": _ids(d["loader"]),
        "hauler": _ids(d["hauler"]),
        "code": text(d["code"].astype("string")).str.extract(r"^(\d{3})", expand=False),
        "remark": text(d["remark"].astype("string")),
    })
    problems = []
    for i, (h, r) in enumerate(zip(d["hour"], out.itertuples(), strict=True), 1):
        if r.slot is None or pd.isna(r.slot):
            hv = "" if h is None or pd.isna(h) else str(h)
            problems.append(f"Remark line {i}: hour '{hv}' is not an hour of the {shift} shift "
                            f"({SLOTS[shift][0]} … {SLOTS[shift][-1]}).")
        if pd.isna(r.loader):
            problems.append(f"Remark line {i}: pick the loader.")
        elif loaders is not None and r.loader not in loaders:
            problems.append(f"Remark line {i}: loader {r.loader} is not in this shift's lines.")
        if pd.isna(r.code) and pd.isna(r.remark):
            problems.append(f"Remark line {i}: give a remark code or a text.")
    out["slot"] = pd.to_numeric(out["slot"], errors="coerce").astype("Int64")
    return out.sort_values(["slot", "loader"]).reset_index(drop=True), problems


def group_remarks(df: pd.DataFrame) -> pd.DataFrame:
    """Per-hour remarks → events: the same loader, hauler, code and text in consecutive hours become one event
    (slot_from..slot_to). `ids` keeps the stored rows of each event (to delete it as a whole)."""
    cols = ["slot_from", "slot_to", "loader", "hauler", "code", "remark", "ids"]
    if df is None or df.empty:
        return pd.DataFrame(columns=cols)
    d = df.assign(id=df["id"] if "id" in df else range(len(df)))
    key = ["loader", "hauler", "code", "remark"]
    d = d.assign(**{k: d[k].astype(object).where(d[k].notna(), None) for k in key})
    d = d.assign(_k=["|".join(map(str, t)) for t in d[key].itertuples(index=False)]).sort_values(["_k", "slot"])
    out = []
    for _, g in d.groupby("_k", sort=False):
        k = tuple(g.iloc[0][key])
        slots = [int(x) for x in g["slot"]]
        ids = g["id"].tolist()
        start = prev = slots[0]
        bag = [ids[0]]
        for s_, i in zip(slots[1:], ids[1:], strict=True):
            if s_ <= prev + 1:
                prev, bag = max(prev, s_), [*bag, i]
                continue
            out.append((start, prev, *k, bag))
            start = prev = s_
            bag = [i]
        out.append((start, prev, *k, bag))
    res = pd.DataFrame(out, columns=cols)
    for c in ("loader", "hauler", "code", "remark"):
        res[c] = res[c].astype(object).where(res[c].notna(), None)
    return res.sort_values(["slot_from", "loader"]).reset_index(drop=True)


def span_label(shift: str, a: int, b: int) -> str:
    """'09-10' for one hour, '09-10 – 11-12' for several."""
    return SLOTS[shift][a - 1] if a == b else f"{SLOTS[shift][a - 1]} – {SLOTS[shift][b - 1]}"


# short words for the TV cells: letters, so a marker can never be read as BCM or tonnes
REMARK_TAGS = {"100": "OK", "101": "SHIFT", "301": "STBY", "302": "RAIN", "303": "SLIP", "304": "BLAST",
               "305": "WAIT", "401": "BD-L", "402": "BD-H", "501": "PM", "502": "FRONT"}
REMARK_GROUP_TAG = {"1": "INFO", "2": "IDLE", "3": "DELAY", "4": "BD", "5": "MAINT"}
REMARK_CATEGORY = {"1": "info", "2": "delay", "3": "delay", "4": "down", "5": "maint"}   # by the first digit


def remark_tag(code) -> str:
    """'302' → 'RAIN'; any other 3-digit code → its group word (4xx → 'BD'); no code (text only) → 'NOTE'."""
    if code is None or pd.isna(code) or not str(code):
        return "NOTE"
    c = str(code)
    return REMARK_TAGS.get(c) or REMARK_GROUP_TAG.get(c[0], "NOTE")


def remark_category(code) -> str:
    """info | delay | down | maint | note: the colour of the marker on the TV."""
    if code is None or pd.isna(code) or not str(code):
        return "note"
    return REMARK_CATEGORY.get(str(code)[0], "note")


def remark_label(code) -> str | None:
    """'302' → '302 - Rain'; other 3-digit codes stay as they are."""
    if code is None or pd.isna(code):
        return None
    return f"{code} - {REMARKS[str(code)]}" if str(code) in REMARKS else str(code)


# ------------------------------------------------------------------ time
def production_hour(now: dt.datetime) -> tuple[dt.date, str, int]:
    """(production date, shift, slot 1..12) for a WIB time. The production day starts at 06:00: the night shift
    after midnight still belongs to the previous date."""
    h = now.hour
    if 6 <= h < 18:
        return now.date(), "DS", h - 6 + 1
    if h >= 18:
        return now.date(), "NS", h - 18 + 1
    return now.date() - dt.timedelta(days=1), "NS", h + 6 + 1


def slot_label(shift: str, slot: int) -> str:
    return SLOTS[shift][slot - 1]


# ------------------------------------------------------------------ master data from 'Link Muatan'
def parse_link_muatan(data: bytes) -> tuple[pd.DataFrame, pd.DataFrame]:
    """(load factors [material, material_group, hauler_model, muatan], targets [unit_id, model, target_per_hour])
    from the 'Link Muatan' sheet of the Mst Hourly workbook."""
    raw = read_workbook(data)
    if "Link Muatan" not in raw:
        raise StructureError(["Sheet 'Link Muatan' was not found."])
    x = raw["Link Muatan"]
    hdr = x.iloc[0]
    if str(hdr.iloc[1]).strip() != "Material":
        raise StructureError(["'Link Muatan': cell B1 must be 'Material' with hauler models to its right."])
    models = {j: str(hdr.iloc[j]).strip() for j in range(2, len(hdr)) if pd.notna(hdr.iloc[j])
              and str(hdr.iloc[j]).strip() and str(hdr.iloc[j]).strip() != "Tanggal"}
    models = {j: v for j, v in models.items() if j < 9}      # matrix block B:I; the unit list starts at K
    lf = []
    for i in range(1, len(x)):
        mat = x.iloc[i, 1]
        if pd.isna(mat) or not str(mat).strip():
            break
        mat = str(mat).strip()
        for j, model in models.items():
            v = pd.to_numeric(x.iloc[i, j], errors="coerce")
            if pd.notna(v) and v > 0:
                lf.append((mat, model, float(v)))
    lf = pd.DataFrame(lf, columns=["material", "hauler_model", "muatan"])
    lf.insert(1, "material_group", material_group(lf["material"]).to_numpy())
    tg = x.iloc[1:, 10:13].copy()
    tg.columns = ["unit_id", "model", "target_per_hour"]
    tg["unit_id"] = _ids(tg["unit_id"])
    tg["target_per_hour"] = num(tg["target_per_hour"])
    tg = tg.dropna(subset=["unit_id", "target_per_hour"])
    tg = tg[tg["unit_id"].str.match(r"^[A-Z]{2,4}\d", na=False)].drop_duplicates("unit_id")
    return lf.reset_index(drop=True), tg.reset_index(drop=True)


# ------------------------------------------------------------------ resolving input rows
@dataclass
class Resolved:
    rows: pd.DataFrame
    problems: list[str] = field(default_factory=list)     # blocking
    warnings: list[str] = field(default_factory=list)


def nrp_of(v) -> str | None:
    """'12345 - Budi S.' (drop-down label) or '12345' → '12345'."""
    if v is None or (isinstance(v, float) and pd.isna(v)) or (not isinstance(v, str) and pd.isna(v)):
        return None
    s = str(v).strip()
    if not s:
        return None
    s = s.split(" - ", 1)[0].strip()
    return s[:-2] if s.endswith(".0") and s[:-2].isdigit() else s


def is_hauler(units: pd.DataFrame) -> pd.Series:
    """Hauling units: type or description mentions 'haul' (some dump trucks are typed 'Supporting Equipment'
    with the description 'Hauling 23 Ton')."""
    txt = units["type"].fillna("") + " " + (units["description"].fillna("") if "description" in units else "")
    return txt.str.contains("haul", case=False)


def load_model(unit_model, lf_models, mapping: dict | None = None) -> str | None:
    """Load-factor model for a population model: the mapping first, then the same name, then the name without the
    unit suffix ('777E-KDP' → '777E', '773E-PRB' → '773E')."""
    if unit_model is None or (isinstance(unit_model, float) and pd.isna(unit_model)):
        return None
    um = str(unit_model).strip().upper()
    known = {str(x).strip().upper(): x for x in lf_models}
    if mapping and um in mapping:
        return mapping[um]
    if um in known:
        return known[um]
    base = um.split("-")[0].strip()
    return known.get(base)


def suggest_load_model(unit_model: str, lf_models) -> str | None:
    """Best guess for the mapping page: the automatic match, else the load class sharing the longest leading part
    of the name (CWE37064R → CWE370Q, 775F-DLS → 775E); at least three characters must match."""
    hit = load_model(unit_model, lf_models)
    if hit:
        return hit
    um = str(unit_model).upper()

    def common(a: str, b: str) -> int:
        n = 0
        for x, y in zip(a, b, strict=False):
            if x != y:
                break
            n += 1
        return n
    best = max(lf_models, key=lambda x: common(um, str(x).upper()), default=None)
    return best if best is not None and common(um, str(best).upper()) >= 3 else None


def expand_to_population(load: pd.DataFrame, pop_models) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Load factors keyed by the models of the unit population. Each population model without its own values takes
    them from the matching general model of Link Muatan ('777E-KDP' ← '777E'), or, failing that, from the closest
    name (to be checked). Returns (load factors incl. the originals, report: model, from, how)."""
    general = sorted(set(load["hauler_model"])) if len(load) else []
    have = {str(x).upper() for x in general}
    rows, report = [load], []
    for pm in sorted({str(x) for x in pop_models if isinstance(x, str) and x.strip()}):
        if pm.upper() in have:
            report.append((pm, pm, "own values"))
            continue
        src, how = load_model(pm, general), "same model"
        if src is None:
            src, how = suggest_load_model(pm, general), "closest name — check"
        if src is None:
            report.append((pm, None, "no match — fill in"))
            continue
        rows.append(load[load["hauler_model"] == src].assign(hauler_model=pm))
        report.append((pm, src, how))
    out = pd.concat(rows, ignore_index=True).drop_duplicates(["material", "hauler_model"])
    return out, pd.DataFrame(report, columns=["model", "values from", "how"])


def resolve(rows: pd.DataFrame, load: pd.DataFrame, targets: pd.DataFrame, units: pd.DataFrame | None = None,
            operators: pd.DataFrame | None = None, model_map: dict | None = None,
            model_targets: pd.DataFrame | None = None,
            hourly_models: pd.DataFrame | None = None, locations: pd.DataFrame | None = None) -> Resolved:
    """Fill hauler model, load, material group, loader model, hourly target and operator names; check every line.

    Target of each excavator (core.prod_target.hourly_target): `targets` = unit overrides, `hourly_models` = the
    site's Hourly Production targets per model, `model_targets` = the Production Data defaults (the fallback,
    marked 'default' in target_source).

    `locations` (core.locations: kind PIT | DISPOSAL, name, material_group, active) of the site: when given, a line
    with trips needs an active PIT and an active disposal of its material. distance_m (horizontal) and dist_v are
    typed per line and kept as typed (a line with trips but no distance only warns).

    rows: loader, loader_nrp, hauler, hauler_nrp, material, pit, disposal, r1..r12, remark_code, remark;
    Legacy lines without a hauler ID may give hauler_model instead. One line = one hauler for one loader and one disposal; the same hauler may
    appear on several lines (another disposal, or its operator changed)."""
    out = rows.copy()
    for c in R:
        out[c] = num(out[c]) if c in out else np.nan
    for c in ("loader", "hauler"):
        out[c] = _ids(out[c]) if c in out else pd.NA
    for c in ("material", "hauler_model", "operator", "hauler_operator", "pit", "disposal", "remark", "remark_code"):
        out[c] = text(out[c].astype("string")) if c in out else pd.NA
    for c in ("loader_nrp", "hauler_nrp"):
        out[c] = out[c].map(nrp_of) if c in out else None
    for c in ("distance_m", "dist_v"):
        out[c] = num(out[c]) if c in out else np.nan
    has_trips = out[R].fillna(0).sum(axis=1) > 0
    legacy = (out["hauler_model"].notna() & out["loader"].notna()) if "hauler_model" in out else False
    truck = out["hauler"].notna() | legacy
    # pre-filled lines left without a truck and without trips are unused, not mistakes
    out = out[truck | has_trips].reset_index(drop=True)
    problems, warnings = [], []

    out["hauler_model"] = out["hauler_model"].mask(out["hauler_model"] == "?")   # Excel lookup found nothing
    models = dict(zip(units["unit_id"], units["model"], strict=True)) if units is not None and len(units) else {}
    out["hauler_model"] = out["hauler"].map(models).fillna(out["hauler_model"]) if len(out) else out["hauler_model"]
    lf = {(str(a).upper(), str(b).upper()): (v, g) for a, b, v, g in
          load[["material", "hauler_model", "muatan", "material_group"]].itertuples(index=False)}
    lf_models = sorted(set(load["hauler_model"])) if len(load) else []
    mapping = {str(k).strip().upper(): v for k, v in (model_map or {}).items()}
    out["muatan"] = np.nan
    out["material_group"] = material_group(out["material"]).to_numpy()
    for i, r in out.iterrows():
        who = next((str(v) for v in (r["hauler"], r["hauler_model"]) if pd.notna(v)), "no hauler")
        line = f"Line {i + 1} ({r['loader'] if pd.notna(r['loader']) else 'no loader'} · {who})"
        if pd.isna(r["loader"]):
            problems.append(f"{line}: loader is empty.")
        if pd.isna(r["hauler"]) and pd.isna(r["hauler_model"]):
            problems.append(f"{line}: hauler is empty.")
            continue
        if pd.isna(r["hauler_model"]):
            problems.append(f"{line}: hauler {r['hauler']} is not in the unit population, so its model and load "
                            "are unknown (Input & upload → Upload Unit Population).")
            continue
        if pd.isna(r["material"]):
            problems.append(f"{line}: material is empty.")
            continue
        lm = load_model(r["hauler_model"], lf_models, mapping)
        if lm is None:
            problems.append(f"{line}: hauler model {r['hauler_model']} has no load per trip; add it in "
                            "Setup → Load factors.")
            continue
        hit = lf.get((str(r["material"]).upper(), str(lm).upper()))
        if hit is None:
            import difflib
            mats = sorted({a for a, b in lf if b == str(lm).upper()})
            near = difflib.get_close_matches(str(r["material"]).upper(), mats, n=1, cutoff=0.75)
            hint = (f" Did you mean '{next(x for x in load['material'] if str(x).upper() == near[0])}'? Pick the "
                    "material from the drop-down, or rename it in Load factors.") if near else                 " Add it in Load factors."
            problems.append(f"{line}: no load factor for {r['material']} × {lm}.{hint}")
            continue
        out.loc[i, ["muatan", "material_group"]] = list(hit)
        bad = out.loc[i, R][(out.loc[i, R] < 0) | (out.loc[i, R] > MAX_TRIPS)].dropna()
        if len(bad):
            problems.append(f"{line}: trips of one hauler in one hour must be between 0 and {MAX_TRIPS}.")
    _locations(out, locations, problems, warnings)

    from core.prod_target import SOURCE_LABEL, hourly_target
    ov = targets
    unit_model = dict(zip(ov["unit_id"], ov["model"], strict=True)) if ov is not None and len(ov) else {}
    out["loader_model"] = out["loader"].map(unit_model)
    if models:
        out["loader_model"] = out["loader_model"].fillna(out["loader"].map(models))
        unknown = sorted(set(out["loader"].dropna()) - set(models))
        if unknown:
            warnings.append(f"Loader not in the unit population: {', '.join(unknown)}.")
    hits = [hourly_target(u, mdl, mat, ov, hourly_models, model_targets)
            for u, mdl, mat in zip(out["loader"], out["loader_model"], out["material"], strict=True)]
    out["target_per_hour"] = pd.array([h[0] for h in hits], dtype="Float64").astype(float)
    out["target_source"] = [h[1] for h in hits]
    no_target = sorted(set(out.loc[out["target_per_hour"].isna(), "loader"].dropna()))
    if no_target:
        warnings.append(f"No target for: {', '.join(no_target)} (Setup → Hourly targets).")
    on_default = sorted(set(out.loc[out["target_source"] == "default", "loader"].dropna()))
    if on_default:
        warnings.append(f"Using the {SOURCE_LABEL['default']} (no hourly target yet): {', '.join(on_default)}.")

    names = dict(zip(operators["nrp"], operators["name"], strict=True)) if operators is not None and len(operators) \
        else {}
    out["operator"] = out["loader_nrp"].map(names).fillna(out["operator"]) if names else out["operator"]
    out["hauler_operator"] = out["hauler_nrp"].map(names).fillna(out["hauler_operator"]) if names \
        else out["hauler_operator"]
    if names:
        unknown = sorted({x for x in (*out["loader_nrp"], *out["hauler_nrp"]) if isinstance(x, str) and x not in names})
        if unknown:
            warnings.append(f"NRP not in the operator master: {', '.join(unknown)} (Setup → Operators).")
    missing = out[(out[R].fillna(0).sum(axis=1) > 0) & (out["hauler_nrp"].isna()) & out["hauler"].notna()]
    if len(missing):
        warnings.append(f"{len(missing)} line(s) with trips but no hauler operator: operator KPIs will miss them.")
    busy = out.melt(id_vars=["hauler", "loader"], value_vars=R).dropna(subset=["value"])
    busy = busy[(busy["value"] > 0) & busy["hauler"].notna()]
    per_hour = busy.groupby(["hauler", "variable"])["value"].sum()
    over = per_hour[per_hour > MAX_TRIPS]
    for (hauler, rc), v in over.items():
        problems.append(f"Hauler {hauler}: {v:,.0f} trips in hour {int(rc[1:])} over all its lines; one hauler can "
                        f"make at most {MAX_TRIPS} trips in an hour.")
    two = busy.groupby(["hauler", "variable"])["loader"].nunique()
    two = sorted({h for (h, _), n in two.items() if n > 1})
    if two:
        warnings.append(f"Same hauler loaded by more than one excavator in the same hour: {', '.join(two)} "
                        "(fine when it moved to another excavator).")
    out["line"] = range(1, len(out) + 1)
    return Resolved(out, problems, warnings)


def _locations(out: pd.DataFrame, locations: pd.DataFrame | None, problems: list[str],
               warnings: list[str]) -> None:
    """Check the PIT and the disposal of each line against the site's master (in place); distances stay as typed."""
    if out.empty:
        return
    has_trips = out[R].fillna(0).sum(axis=1) > 0
    no_dist = int((has_trips & (out["distance_m"].isna() | out["dist_v"].isna())).sum())
    if no_dist:
        warnings.append(f"{no_dist} line(s) with trips but no H or V distance: the engineering checker types them.")
    if locations is None or locations.empty:
        return
    act = locations[locations["active"].astype(bool)] if "active" in locations else locations
    known = {}                                    # (kind, NAME) → (name as listed, {materials})
    for k, n, g in zip(act["kind"], act["name"], act["material_group"], strict=True):
        known.setdefault((k, str(n).upper()), (n, set()))[1].add(g)
    where = "(Setup → PIT & disposals)"
    for i, r in out.iterrows():
        line = f"Line {i + 1} ({r['loader'] if pd.notna(r['loader']) else 'no loader'} · "                f"{r['hauler'] if pd.notna(r['hauler']) else 'no hauler'})"
        for col, kind, label in (("pit", "PIT", "PIT"), ("disposal", "DISPOSAL", "disposal")):
            v = r[col]
            if pd.isna(v):
                if has_trips[i]:
                    problems.append(f"{line}: pick the {label}.")
                continue
            hit = known.get((kind, str(v).upper()))
            if hit is None:
                problems.append(f"{line}: {label} {v} is not an active {label} of this site {where}.")
                continue
            out.loc[i, col] = hit[0]
            if r["material_group"] in ("OB", "CG") and r["material_group"] not in hit[1]:
                problems.append(f"{line}: {hit[0]} is a {'/'.join(sorted(hit[1]))} {label} but the material is "
                                f"{r['material']}.")


def to_long(rows: pd.DataFrame) -> pd.DataFrame:
    """Rows (with shift, date, site) → one record per line × hour: hour_slot, slot, rit, volume."""
    if rows.empty:
        return pd.DataFrame(columns=["site", "date", "shift", "loader", "hauler", "material_group", "hour_slot",
                                     "slot", "rit", "volume", "distance_m", "dist_v", "hauler_model", "operator",
                                     "hauler_operator", "loader_nrp", "hauler_nrp", "muatan", "target_per_hour"])
    keep = [c for c in rows.columns if c not in R]
    long = rows.melt(id_vars=keep, value_vars=R, var_name="rc", value_name="rit")
    long["slot"] = long["rc"].str[1:].astype(int)
    long["hour_slot"] = [SLOTS[s][k - 1] for s, k in zip(long["shift"], long["slot"], strict=True)]
    long["rit"] = long["rit"].fillna(0.0)
    long["volume"] = long["rit"] * long["muatan"]
    return long.drop(columns="rc")


# ------------------------------------------------------------------ official ritase from the cutover date on
RITASE_COLS = ["site", "site_hauler", "date", "hour_slot", "shift", "hauler", "hauler_model", "muatan", "loader",
               "loader_model", "material", "material_group", "pit", "disposal", "dist_v", "dist_h", "rit", "volume"]


def as_ritase(long: pd.DataFrame) -> pd.DataFrame:
    """Shift lines per hour (to_long) → the columns of the Production Data ritase table, hours with trips only.
    The hauler counts to the loader's site, as in Production Data."""
    if long is None or long.empty:
        return pd.DataFrame(columns=RITASE_COLS)
    d = long[long["rit"] > 0]
    out = d.assign(site_hauler=d["site"], dist_h=d.get("distance_m", np.nan))
    return out.reindex(columns=RITASE_COLS).reset_index(drop=True)


def official_ritase(imported: pd.DataFrame, hourly: pd.DataFrame, cutover: dt.date | None) -> pd.DataFrame:
    """Ritase used by dashboards and TVs: Production Data before the cutover date, Hourly Production from it on (all
    saved shifts, approved or not). No cutover → Production Data only."""
    if cutover is None:
        return imported
    before = imported[imported["date"] < cutover] if len(imported) else imported
    after = hourly[hourly["date"] >= cutover] if hourly is not None and len(hourly) else None
    parts = [p for p in (before, after) if p is not None and len(p)]
    if not parts:
        return imported.iloc[:0] if len(imported.columns) else pd.DataFrame(columns=RITASE_COLS)
    return pd.concat(parts, ignore_index=True)


# ------------------------------------------------------------------ Excel template (one shift per file)
def _label(nrp, names: dict) -> str | None:
    nrp = nrp_of(nrp)
    return None if not nrp else (f"{nrp} - {names[nrp]}" if nrp in names else nrp)


def build_template(site: str, date: dt.date, shift: str, load: pd.DataFrame, targets: pd.DataFrame,
                   lines: pd.DataFrame | None = None, coordinator: str = "", units: pd.DataFrame | None = None,
                   operators: pd.DataFrame | None = None, remarks: pd.DataFrame | None = None,
                   locations: pd.DataFrame | None = None) -> bytes:
    """Per-shift input workbook, one row per hauler and disposal. `lines` pre-fills loader, hauler, operators, pit,
    disposal and distances (e.g. from the previous shift) so the data officer only types trips. Drop-downs list the
    site's loaders, haulers, operators and the active PITs and disposals (`locations`)."""
    from openpyxl.utils import get_column_letter

    from core import dataprod
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws["A1"] = f"PRISMA · {DATASET} · {shift} · one row per hauler"
    ws["A1"].font = Font(bold=True, size=14)
    for r, (k, v) in enumerate([("Site", site), ("Date", date), ("Shift", shift), ("Shift boss", coordinator)],
                               start=2):
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
    ws["B3"].number_format = "yyyy-mm-dd"
    for r, tip in enumerate(["How to fill: one row per truck and disposal. Pick Loader, Operator, Material, Hauler "
                             "ID, PIT and Disposal from the drop-downs, then type the trips per hour.",
                             "Truck model and load are added automatically. H and V distance (m) are typed per row "
                             "by the engineering checker. Rows without a Hauler ID and without trips are ignored.",
                             "Truck went to two disposals, or its operator changed? Add a second row for the same "
                             "truck."], start=2):
        ws.cell(r, 4, tip).font = Font(italic=True, color="55595F")
    heads = INPUT_COLS + SLOTS[shift]
    widths = {"Loader": 11, "Hauler ID": 11, "Operator": 24, "Hauler operator": 24, "Material": 18, "Remark": 28,
              "Disposal": 22, "PIT": 14, "H distance (m)": 13, "V distance (m)": 13}
    for j, h in enumerate(heads, start=1):
        c = ws.cell(HEADER_ROW, j, h)
        c.fill, c.font = dataprod.HEAD_FILL, dataprod.HEAD_FONT
        ws.column_dimensions[c.column_letter].width = 7 if h in SLOTS[shift] else widths.get(h, 14)
    names = dict(zip(operators["nrp"], operators["name"], strict=True)) if operators is not None and len(operators) \
        else {}
    u = units if units is not None else pd.DataFrame(columns=["unit_id", "type", "site"])
    u = u[u["site"] == site] if len(u) else u
    hu = u[is_hauler(u)].sort_values("unit_id") if len(u) else u
    haulers = hu["unit_id"].tolist() if len(hu) else []
    loaders = sorted(set(targets["unit_id"]) if len(targets) else set())
    lists = wb.create_sheet("Lists")
    from core.locations import names as loc_names
    columns = {"A": sorted(load["material"].unique()) if len(load) else [],
               "B": [f"{k} - {v}" for k, v in REMARKS.items()], "C": loaders, "D": haulers,
               "E": [f"{n} - {nm}" for n, nm in sorted(names.items(), key=lambda x: x[1])],
               "G": loc_names(locations, "DISPOSAL"),               # F holds the hour labels (Remarks sheet)
               "H": loc_names(locations, "PIT")}
    for col, vals in columns.items():
        for i, v in enumerate(vals, start=1):
            lists[f"{col}{i}"] = v
    lists.sheet_state = "hidden"
    last = HEADER_ROW + 400
    pos = {h: get_column_letter(j) for j, h in enumerate(heads, start=1)}
    for head, src in (("Material", "A"), ("Loader", "C"), ("Hauler ID", "D"),
                      ("Operator", "E"), ("Hauler operator", "E"), ("Disposal", "G"), ("PIT", "H")):
        n = len(columns[src])
        if n:
            dv = DataValidation(type="list", formula1=f"=Lists!${src}$1:${src}${n}", allow_blank=True,
                                errorStyle="warning")
            dv.add(f"{pos[head]}{HEADER_ROW + 1}:{pos[head]}{last}")
            ws.add_data_validation(dv)
    first, lastc = pos[SLOTS[shift][0]], pos[SLOTS[shift][-1]]
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="20", allow_blank=True)
    dv.error, dv.errorTitle = "Trips of one hauler per hour: whole number 0–20", "Trips"
    dv.add(f"{first}{HEADER_ROW + 1}:{lastc}{last}")
    ws.add_data_validation(dv)
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        errorStyle="warning")
    dv.error, dv.errorTitle = "Distance in metres, usually 0 or more", "Distance"
    for h in ("H distance (m)", "V distance (m)"):
        dv.add(f"{pos[h]}{HEADER_ROW + 1}:{pos[h]}{last}")
    ws.add_data_validation(dv)
    recs = lines.to_dict("records") if lines is not None else []
    for i, r in enumerate(recs, start=HEADER_ROW + 1):
        name = r.get("operator") if isinstance(r.get("operator"), str) else None
        vals = {"Loader": r.get("loader"), "Operator": _label(r.get("loader_nrp"), names) or name,
                "Material": r.get("material"), "Hauler ID": r.get("hauler"),
                "Hauler operator": _label(r.get("hauler_nrp"), names), "PIT": r.get("pit"), "Disposal": r.get("disposal"),
                "H distance (m)": r.get("distance_m"), "V distance (m)": r.get("dist_v")}
        for h, v in vals.items():
            ws[f"{pos[h]}{i}"] = None if v is None or (isinstance(v, float) and pd.isna(v)) else v
    ws.freeze_panes = ws.cell(HEADER_ROW + 1, 5)          # loader, operator, material and truck stay visible
    _remark_sheet(wb, shift, columns, remarks)
    dataprod._meta(wb, "hourly", {"dataset": DATASET, "site": site, "date": date.isoformat(), "shift": shift, "layout": "per_hauler"})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _remark_sheet(wb, shift: str, lists: dict, remarks: pd.DataFrame | None) -> None:
    """Sheet 'Remarks': one line per event, only when something happened (rain at 09-10, breakdown at 13-14…)."""
    from core import dataprod
    rs = wb.create_sheet(REMARK_SHEET, 1)
    rs["A1"] = "Remarks per hour: one line per event. Hour, loader and a remark code (or text) are needed."
    rs["A1"].font = Font(italic=True, color="55595F")
    for j, h in enumerate(REMARK_HEADS, start=1):
        c = rs.cell(2, j, h)
        c.fill, c.font = dataprod.HEAD_FILL, dataprod.HEAD_FONT
        rs.column_dimensions[c.column_letter].width = {"Remark code": 30, "Remark": 40}.get(h, 12)
    lst = wb["Lists"]
    for i, lab in enumerate(SLOTS[shift], start=1):
        lst[f"F{i}"] = lab
    for col, src, n in (("A", "F", 12), ("B", "C", len(lists["C"])), ("C", "D", len(lists["D"])),
                        ("D", "B", len(lists["B"]))):
        if n:
            dv = DataValidation(type="list", formula1=f"=Lists!${src}$1:${src}${n}", allow_blank=True,
                                errorStyle="warning")
            dv.add(f"{col}3:{col}300")
            rs.add_data_validation(dv)
    for i, r in enumerate((remarks.to_dict("records") if remarks is not None else []), start=3):
        k = r.get("slot")
        rs.cell(i, 1, SLOTS[shift][int(k) - 1] if k is not None and not pd.isna(k) else None)
        rs.cell(i, 2, r.get("loader"))
        rs.cell(i, 3, None if pd.isna(r.get("hauler")) else r.get("hauler"))
        rs.cell(i, 4, remark_label(r.get("code")))
        rs.cell(i, 5, None if pd.isna(r.get("remark")) else r.get("remark"))
    rs.freeze_panes = "A3"


@dataclass
class HourlyFile:
    site: str
    date: dt.date
    shift: str
    coordinator: str
    rows: pd.DataFrame
    remarks: pd.DataFrame = field(default_factory=lambda: pd.DataFrame(columns=["hour", "loader", "hauler", "code",
                                                                              "remark"]))


def file_name(site: str, date: dt.date, shift: str) -> str:
    """Hourly_Production_WBK-BAU_2026-09-27_DS.xlsx"""
    return f"{file_stem(DATASET)}_{site}_{date:%Y-%m-%d}_{shift}.xlsx"


def parse_template(data: bytes) -> HourlyFile:
    raw = read_workbook(data)
    name = next((n for n in SHEET_NAMES if n in raw), None)
    if name is None:
        raise StructureError([f"Sheet '{SHEET}' was not found: use the {DATASET} template of this site and shift."])
    x = raw[name]
    head = {str(x.iloc[i, 0]).strip(): x.iloc[i, 1] for i in range(1, 5) if i < len(x) and pd.notna(x.iloc[i, 0])}
    problems = []
    site = str(head.get("Site") or "").strip()
    shift = str(head.get("Shift") or "").strip().upper()
    date = excel_date(pd.Series([head.get("Date")])).iloc[0]
    if not site:
        problems.append("Cell B2 (Site) is empty.")
    if shift not in SHIFTS:
        problems.append("Cell B4 (Shift) must be DS or NS.")
    if pd.isna(date):
        problems.append("Cell B3 (Date) must be a date.")
    if problems:
        raise StructureError(problems)
    hdr = [str(v).strip() if pd.notna(v) else "" for v in x.iloc[HEADER_ROW - 1]]
    expected = INPUT_COLS + SLOTS[shift]
    v1 = INPUT_COLS_V1 + SLOTS[shift]                    # templates before destinations: pit/distance typed
    v2 = INPUT_COLS_V2 + SLOTS[shift]                    # destination without typed distances
    v3 = INPUT_COLS_V3 + SLOTS[shift]                    # destination with typed distances, no PIT
    if hdr[:len(expected)] == expected:
        names = ["loader", "loader_nrp", "material", "hauler", "hauler_nrp", "pit", "disposal", "distance_m", "dist_v",
                 *R]
    elif hdr[:len(v3)] == v3:
        expected = v3
        names = ["loader", "loader_nrp", "material", "hauler", "hauler_nrp", "disposal", "distance_m", "dist_v", *R]
    elif hdr[:len(v2)] == v2:
        expected = v2
        names = ["loader", "loader_nrp", "material", "hauler", "hauler_nrp", "disposal", *R]
    elif hdr[:len(v1)] == v1:
        expected = v1
        names = ["loader", "loader_nrp", "material", "hauler", "hauler_nrp", "pit", "disposal", "distance_m", *R]
    else:
        raise StructureError([f"Row {HEADER_ROW} must hold the headers: {', '.join(expected)} "
                              f"(the hour columns follow the shift in B4). Download the current template."])
    body = x.iloc[HEADER_ROW:, :len(expected)].copy()   # old templates: Remark code / Remark after the hours, ignored
    body.columns = names
    body = body.dropna(how="all").assign(remark_code=None, remark=None)
    remarks = pd.DataFrame(columns=["hour", "loader", "hauler", "code", "remark"])
    rs = raw.get(REMARK_SHEET)
    if rs is not None and len(rs) > 2:
        rb = rs.iloc[2:, :len(REMARK_HEADS)].copy()
        rb.columns = ["hour", "loader", "hauler", "code", "remark"][:rb.shape[1]]
        remarks = rb.reindex(columns=["hour", "loader", "hauler", "code", "remark"]).dropna(how="all")
    boss = head.get("Shift boss", head.get("Coordinator"))    # older templates say 'Coordinator'
    boss = "" if boss is None or (isinstance(boss, float) and pd.isna(boss)) or str(boss).strip().lower() == "nan" \
        else str(boss).strip()
    return HourlyFile(site, date.date(), shift, boss, body.reset_index(drop=True), remarks.reset_index(drop=True))


# ------------------------------------------------------------------ legacy 'Mst Hourly' shift sheets
def parse_mst_shift(sheet: pd.DataFrame) -> pd.DataFrame:
    """Rows from a DS/NS sheet of the old Mst Hourly workbook (blocks 'Ritasi OB' and 'Ritasi CG').

    Columns used: D excavator, F operator, G material, H hauler model, M..X trips per hour. Lines without a
    material are template filler and are skipped."""
    out = []
    loader, operator = None, None
    for i in range(len(sheet)):
        r = sheet.iloc[i]
        exca = r.iloc[3] if len(r) > 3 else None
        mat = r.iloc[6] if len(r) > 6 else None
        if isinstance(exca, str) and exca.strip() not in ("", "Exca", "0", loader):
            loader, operator = exca.strip(), None
        if isinstance(r.iloc[5], str) and r.iloc[5].strip() and r.iloc[5].strip() != "Operator":
            operator = r.iloc[5].strip()
        if not (isinstance(mat, str) and mat.strip()[:2].upper() in ("OB", "CG")) or loader is None:
            continue
        trips = [pd.to_numeric(v, errors="coerce") for v in r.iloc[12:24].tolist()]
        out.append({"loader": loader, "operator": operator, "loader_nrp": None, "hauler": None, "hauler_nrp": None,
                    "material": mat.strip(),
                    "hauler_model": str(r.iloc[7]).strip() if pd.notna(r.iloc[7]) else None,
                    "pit": None, "disposal": None, "distance_m": None,
                    **{f"r{k}": trips[k - 1] for k in range(1, 13)}, "remark_code": None, "remark": None})
    return pd.DataFrame(out, columns=["loader", "operator", "loader_nrp", "hauler", "hauler_nrp", "material",
                                      "hauler_model", "pit", "disposal", "distance_m", *R, "remark_code", "remark"])
