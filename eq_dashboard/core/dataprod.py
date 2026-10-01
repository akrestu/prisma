"""Production Data workbooks: the blank template, the export of stored data, and the file-name convention.

Both the template and the export are written from core.validate.SHEETS, so an exported file can be corrected in
Excel and uploaded again (round trip), and a filled template always passes the structure check.
"""
from __future__ import annotations

import datetime as dt
import io
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.config import now_wib
from core.validate import (
    DATA_SHEETS,
    DATASET,
    EVENTS,
    FUEL,
    HOUR_SLOTS,
    META_SHEET,
    RECEIPTS,
    TEMPLATE_VERSION,
    TRIPS,
    WEIGHBRIDGE,
    Sheet,
    file_stem,
)

# Production_Data_2026-09 (current) or Data_Prod_2026-09 (older files)
NAME_RE = re.compile(r"(?:production[_ -]?data|data[_ -]?prod)[_ -]?(\d{4})[-_]?(\d{2})", re.I)
MAX_ROWS = 100_000          # validation / formatting range per sheet
HEAD_FILL = PatternFill("solid", fgColor="F2C230")  # PRISMA hi-vis yellow
HEAD_FONT = Font(bold=True, color="141517")
TITLE_FONT = Font(bold=True, size=13)
FORMATS = {"date": "yyyy-mm-dd", "time": "hh:mm", "datetime": "yyyy-mm-dd hh:mm", "number": "#,##0.00",
           "int": "0"}
SHIFT_FUEL = {"DS": "I", "NS": "II"}


# ------------------------------------------------------------------ file names
def file_name(month: dt.date, site: str | None = None, ext: str = "xlsx") -> str:
    """Production_Data_2026-09.xlsx / Production_Data_2026-09_WBK-MAS.xlsx"""
    return f"{file_stem(DATASET)}_{month:%Y-%m}" + (f"_{site}" if site else "") + f".{ext}"


def month_from_name(name: str) -> dt.date | None:
    m = NAME_RE.search(name or "")
    if not m:
        return None
    y, mo = int(m.group(1)), int(m.group(2))
    return dt.date(y, mo, 1) if 1 <= mo <= 12 else None


def name_check(name: str, data_month: dt.date) -> str | None:
    """A warning when the file name does not follow the convention or names another month (never blocks)."""
    named = month_from_name(name)
    if named is None:
        return (f"File name '{name}' does not follow the convention {file_name(data_month, ext='xlsb')}. "
                "The upload still works; renaming keeps files easy to find.")
    if named != data_month:
        return (f"The file name says {named:%B %Y} but the data is for {data_month:%B %Y}. "
                "Check that you picked the right file.")
    return None


# ------------------------------------------------------------------ workbook building blocks
def _sheet(wb: Workbook, sh: Sheet, rows: list[list] | None = None, lists: dict[str, str] | None = None):
    ws = wb.create_sheet(sh.name)
    hdr = sh.header_row + 1                     # 1-based Excel row of the column names
    if sh.header_row:
        ws.cell(1, 1, f"{sh.name} — {sh.purpose}").font = TITLE_FONT
    for j, c in enumerate(sh.cols, start=1):
        cell = ws.cell(hdr, j, c.name)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        note = c.desc + (f"\nExample: {c.example}" if c.example is not None else "")
        cell.comment = Comment(note, "PRISMA", width=260, height=90)
        letter = get_column_letter(j)
        width = 6 if c.name in HOUR_SLOTS else max(11, min(34, len(c.name) + 4))
        if c.kind in ("datetime",):
            width = 18
        ws.column_dimensions[letter].width = width
        rng = f"{letter}{hdr + 1}:{letter}{hdr + MAX_ROWS}"
        dv = _validation(c, sh, lists or {})
        if dv is not None:
            dv.add(rng)
            ws.add_data_validation(dv)
    for r in rows or []:
        ws.append(r)
    for j, c in enumerate(sh.cols, start=1):   # number formats on the data rows that exist
        fmt = FORMATS.get(c.kind)
        if fmt and rows:
            for (cell,) in ws.iter_rows(min_row=hdr + 1, max_row=hdr + len(rows), min_col=j, max_col=j):
                cell.number_format = fmt
    ws.freeze_panes = ws.cell(hdr + 1, 1)
    ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(sh.cols))}{hdr + max(len(rows or []), 1)}"
    return ws


def _validation(c, sh: Sheet, lists: dict[str, str]) -> DataValidation | None:
    if c.kind == "list":
        src = lists.get("Site") if c.name == "Site" else None
        formula = src or ('"' + ",".join(c.choices) + '"' if c.choices else None)
        if not formula:
            return None
        strict = c.name == "Status"
        dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                            errorStyle="stop" if strict else "warning")
        dv.error = f"Use one of: {', '.join(c.choices) or 'the listed sites'}"
        dv.errorTitle = f"{c.name}"
    elif c.kind == "number":
        dv = DataValidation(type="decimal", operator="between", formula1="0",
                            formula2="24" if c.name == "Total Jam" else "10000000", allow_blank=True)
        dv.error, dv.errorTitle = ("Hours between 0 and 24" if c.name == "Total Jam" else "Enter a number ≥ 0"), c.name
    elif c.kind == "int":
        dv = DataValidation(type="whole", operator="between", formula1="0", formula2="100", allow_blank=True)
        dv.error, dv.errorTitle = "Whole number of trips (0–100)", c.name
    elif c.kind in ("date", "datetime"):
        dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)  # > 2000-01-01
        dv.error, dv.errorTitle = "Enter a real Excel date, not text", c.name
    elif c.kind == "time":
        dv = DataValidation(type="time", operator="between", formula1="0", formula2="0.999988", allow_blank=True)
        dv.error, dv.errorTitle = "Enter a time such as 06:00", c.name
    else:
        return None
    dv.prompt, dv.promptTitle = c.desc[:250], c.name[:32]
    dv.showErrorMessage = dv.showInputMessage = True
    return dv


def _readme(wb: Workbook, title: str, lines: list[str], sheets: tuple[Sheet, ...] = DATA_SHEETS) -> None:
    ws = wb.active
    ws.title = "README"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16)
    r = 3
    for line in lines:
        ws.cell(r, 1, line)
        r += 1
    r += 1
    heads = ["Sheet", "Column", "Type", "Description", "Example"]
    for j, h in enumerate(heads, start=1):
        cell = ws.cell(r, j, h)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    for sh in sheets:
        for c in sh.cols:
            if c.name in HOUR_SLOTS[1:]:
                continue
            r += 1
            name = "06-07 … 05-06 (24 columns)" if c.name == HOUR_SLOTS[0] else c.name
            desc = "Trips per production hour" if c.name == HOUR_SLOTS[0] else c.desc
            for j, v in enumerate([sh.name, name, c.kind, desc, "" if c.example is None else str(c.example)], 1):
                ws.cell(r, j, v)
    for letter, w in zip("ABCDE", (18, 26, 10, 80, 22), strict=True):
        ws.column_dimensions[letter].width = w


def _meta(wb: Workbook, kind: str, extra: dict[str, str], version: int | None = None) -> None:
    """Hidden key/value sheet. `version` is the layout version of the workbook kind (Production Data: its
    TEMPLATE_VERSION, read by core.validate; other workbooks: 1 unless given)."""
    ws = wb.create_sheet(META_SHEET)
    if version is None:
        version = TEMPLATE_VERSION if extra.get("dataset", DATASET) == DATASET else 1
    rows = {"dataset": DATASET, "template_version": str(version), "kind": kind,
            "generated_at": now_wib().strftime("%Y-%m-%d %H:%M WIB"), "app": "PRISMA", **extra}
    for k, v in rows.items():
        ws.append([k, v])
    ws.sheet_state = "hidden"


def _lists(wb: Workbook, sites: list[str]) -> dict[str, str]:
    ws = wb.create_sheet("Lists")
    ws["A1"] = "Site"
    for i, s in enumerate(sites, start=2):
        ws.cell(i, 1, s)
    ws.sheet_state = "hidden"
    return {"Site": f"=Lists!$A$2:$A${max(len(sites), 1) + 1}"} if sites else {}


def _bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


RULES = [
    "One month per file (month to date is fine), or several months such as a whole year: the app splits it by month "
    "and each site × month is approved on its own. Upload a month again to replace it; old versions are kept.",
    f"Do not rename sheets or column headers. Headers are in row 1, except '{TRIPS}' where they are in row 2. "
    "Files with the old sheet names (Eq.Event, Ritasi Unit, Data Timbangan, Fuel Consume, Fuel Receipt) still work.",
    "Units and their sites come from the separate Unit Population workbook (Input & upload → Unit Population).",
    "Extra columns are allowed and ignored. Empty rows are skipped.",
    "Dates must be real Excel dates, times real Excel times (hover a header to see its description and an example).",
    f"Save as .xlsx or .xlsb and name it {file_stem(DATASET)}_YYYY-MM.xlsx (one month, e.g. "
    f"{file_stem(DATASET)}_2026-09.xlsx) or {file_stem(DATASET)}_YYYY.xlsx (a whole year).",
    f"Upload it in PRISMA → Input & upload → {DATASET} → Import. Structure and data quality are checked before "
    "anything is saved. Small corrections can also be made in the Edit tab without Excel.",
]


# ------------------------------------------------------------------ template
def build_template(sites: list[str]) -> bytes:
    """Blank Production Data workbook. Units and their sites come from the separate Unit Population workbook."""
    wb = Workbook()
    _readme(wb, f"PRISMA · {DATASET} template", [
        "Production & Reliability Information System for Mining Analytics — production data workbook: one month or "
        "a whole year.", "", *RULES])
    lists = _lists(wb, sites)
    for sh in DATA_SHEETS:
        _sheet(wb, sh, None, lists)
    wb.move_sheet("Lists", offset=len(wb.sheetnames))
    _meta(wb, "template", {"sites": ",".join(sites)})
    return _bytes(wb)


def _unit_rows(u: pd.DataFrame) -> list[list]:
    u = u.sort_values(["site", "type", "unit_id"], na_position="last")
    return [[_v(r.type), _v(r.description), r.unit_id, _v(r.model), _v(r.manufacturer), _v(r.site)]
            for r in u.itertuples()]


def _v(x):
    if x is None:
        return None
    try:
        if pd.isna(x):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(x, np.generic):
        return x.item()
    if isinstance(x, pd.Timestamp):
        return x.to_pydatetime(warn=False)
    return x


def _time(frac) -> dt.time | None:
    if frac is None or pd.isna(frac):
        return None
    secs = round(float(frac) % 1.0 * 86400) % 86400
    return dt.time(secs // 3600, secs % 3600 // 60, secs % 60)


# ------------------------------------------------------------------ export (stored data → Production Data workbook)
def export_rows(t: dict[str, pd.DataFrame]) -> dict[str, list[list]]:
    """`t` holds stored tables keyed like repo.EXPORT_TABLES: units, events, ritase, coal, fuel, receipt.
    Returns the rows of every data sheet in template column order; they re-import to the same numbers."""
    return {
        EVENTS: _event_rows(t["events"]),
        TRIPS: _ritase_rows(t["ritase"]),
        WEIGHBRIDGE: [[_v(r.date), _v(r.ticket_id), _v(r.supplier), _v(r.product), _v(r.time_in), _v(r.time_out),
                       _v(r.loader), _v(r.shift), _v(r.dt_unit), _v(r.dist_h), _v(r.ton), _v(r.dist_v)]
                      for r in t["coal"].sort_values(["date", "time_in"]).itertuples()],
        FUEL: [[_v(r.date), SHIFT_FUEL.get(r.shift, _v(r.shift)), _v(r.model), _v(r.unit_id), _time(r.time),
                _v(r.liters)] for r in t["fuel"].sort_values(["date", "unit_id"]).itertuples()],
        RECEIPTS: [[_v(r.date), SHIFT_FUEL.get(r.shift, _v(r.shift)), _v(r.vendor), _v(r.operator), _v(r.unit),
                    _v(r.dn_no), _v(r.liters)] for r in t["receipt"].sort_values("date").itertuples()],
    }


def export_frames(t: dict[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
    """The same rows as DataFrames with the template's column names (the manual-edit grids)."""
    return {sh.name: pd.DataFrame(export_rows(t)[sh.name] or None, columns=[c.name for c in sh.cols])
            for sh in DATA_SHEETS}


def workbook(rows: dict[str, list[list]], month: dt.date, sites: list[str], kind: str, note: str) -> bytes:
    """A Production Data workbook from rows per sheet (export and manual edit share it, so both re-import exactly
    like an uploaded file)."""
    wb = Workbook()
    _readme(wb, f"PRISMA · {DATASET} {kind} · {month:%B %Y} · {', '.join(sites)}", [
        note, "Stoppages, weeks and data quality are recalculated on upload, so they are not part of this file.", "",
        *RULES])
    lists = _lists(wb, sites)
    for sh in DATA_SHEETS:
        _sheet(wb, sh, rows.get(sh.name), lists)
    wb.move_sheet("Lists", offset=len(wb.sheetnames))
    _meta(wb, kind, {"month": f"{month:%Y-%m}", "sites": ",".join(sites)})
    return _bytes(wb)


def export_workbook(t: dict[str, pd.DataFrame], month: dt.date, sites: list[str]) -> bytes:
    return workbook(export_rows(t), month, sites, "export",
                    "Exported from PRISMA (PUBLISHED data). Edit in Excel and upload again to create a new version.")


def frames_workbook(frames: dict[str, pd.DataFrame], month: dt.date, sites: list[str]) -> bytes:
    """Edited grids (export_frames layout) → workbook bytes for the normal import pipeline."""
    rows = {}
    for sh in DATA_SHEETS:
        df = frames[sh.name].dropna(how="all")
        rows[sh.name] = [[_v(x) for x in r] for r in df[[c.name for c in sh.cols]].itertuples(index=False)]
    return workbook(rows, month, sites, "manual_edit", "Made in PRISMA from the manual edit grids.")


def _event_rows(ev: pd.DataFrame) -> list[list]:
    if ev.empty:
        return []
    ev = ev.sort_values(["unit_id", "seq"]) if "seq" in ev.columns else ev
    out = []
    for r in ev.itertuples():
        code = r.reason_code
        reason = (f"{int(code):03d} {r.reason_text or ''}".strip() if code is not None and not pd.isna(code)
                  else _v(r.reason_text))
        out.append([_v(r.date), _v(r.shift), r.unit_id, _v(r.operator), _time(r.time_start), _time(r.time_end),
                    _v(r.hours), _v(r.hm_start), _v(r.hm_end), _v(r.status), reason])
    return out


RIT_KEYS = ["date", "hauler", "hauler_model", "muatan", "loader", "loader_model", "material", "pit", "disposal",
            "dist_v", "dist_h"]


def _ritase_rows(rit: pd.DataFrame) -> list[list]:
    if rit.empty:
        return []
    k = rit[RIT_KEYS].astype(object).where(rit[RIT_KEYS].notna(), "§")   # keep rows with empty keys together
    wide = (rit.assign(**{c: k[c] for c in RIT_KEYS})
               .pivot_table(index=RIT_KEYS, columns="hour_slot", values="rit", aggfunc="sum", fill_value=0)
               .reindex(columns=HOUR_SLOTS, fill_value=0).reset_index())
    rows = []
    for r in wide.itertuples(index=False):
        vals = [None if v == "§" else _v(v) for v in r[:len(RIT_KEYS)]]
        hours = [int(v) if float(v).is_integer() else float(v) for v in r[len(RIT_KEYS):]]
        rows.append(vals + [h if h else None for h in hours])
    return rows
