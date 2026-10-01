"""The three datasets: consistent names, old sheet names still read, manual edit round trip, deleting per dataset."""
from __future__ import annotations

import datetime as dt
import io

import pandas as pd
import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from core import dataprod, population
from core import hourly as H
from core import ingest as ing
from core.io import read_workbook
from core.parse import parse_data_prod
from core.validate import DATA_SHEETS, LEGACY_SHEETS, OLD_NAME, validate
from db import models as m
from db import repo


def test_old_sheet_names_are_still_read():
    raw = read_workbook(dataprod.build_template(["WBK-MAS"]))
    old = {OLD_NAME.get(k, k): v for k, v in raw.items()}
    assert "Eq.Event" in old and "Equipment Events" not in old
    frames = validate(old)
    assert {sh.name for sh in DATA_SHEETS} <= set(frames)
    assert set(LEGACY_SHEETS) == {"Populasi Unit", "Eq.Event", "Ritasi Unit", "Data Timbangan", "Fuel Consume",
                                  "Fuel Receipt"}


def test_dataset_file_names():
    assert population.file_name(dt.date(2026, 9, 1)) == "Unit_Population_2026-09-01.xlsx"
    assert H.file_name("WBK-BAU", dt.date(2026, 9, 27), "DS") == "Hourly_Production_WBK-BAU_2026-09-27_DS.xlsx"
    raw = read_workbook(population.build_template(["WBK-MAS"]))
    assert "Unit Population" in raw


def test_manual_edit_round_trip_of_one_site(parsed):
    """export_frames → (grid) → frames_workbook re-imports to the same numbers for that site."""
    p, site = parsed, "WBK-BAU"
    t = {"units": p.units, "events": p.events[p.events["site"] == site],
         "ritase": p.ritase[p.ritase["site"] == site],
         "coal": p.coal[(p.coal["site"] == site) & ~p.coal["cancelled"]],
         "fuel": p.fuel[p.fuel["site"] == site], "receipt": p.receipt[p.receipt["site"] == site]}
    frames = dataprod.export_frames(t)
    assert list(frames) == [sh.name for sh in DATA_SHEETS]
    q = parse_data_prod(dataprod.frames_workbook(frames, p.month, [site]), population=p.units)
    assert len(q.events[q.events["site"] == site]) == len(t["events"])
    assert q.ritase.loc[q.ritase["site"] == site, "rit"].sum() == t["ritase"]["rit"].sum()
    # an edit: drop one event row → one row fewer after the import pipeline
    frames["Equipment Events"] = frames["Equipment Events"].iloc[1:]
    r = parse_data_prod(dataprod.frames_workbook(frames, p.month, [site]), population=p.units)
    assert len(r.events[r.events["site"] == site]) == len(t["events"]) - 1


def test_ingest_can_limit_the_versions_it_creates(db_session, sample_bytes):
    s = db_session
    _, ups = ing.ingest(s, sample_bytes, "edit.xlsb", username="t", sites=["WBK-BAU"])
    assert [us.site_code for us in ups] == ["WBK-BAU"]


def test_delete_production_per_site_then_the_rest(db_session, sample_bytes):
    s = db_session
    up, _ = ing.ingest(s, sample_bytes, "Eq.Event.xlsb", username="t")
    s.commit()
    n_bau = s.scalar(select(func.count()).select_from(m.FactEvent).where(m.FactEvent.site == "WBK-BAU"))
    preview = repo.delete_production(s, ["WBK-BAU"], dry_run=True)
    assert preview["versions"] == 1 and preview["fact_event"] == n_bau > 0
    assert s.scalar(select(func.count()).select_from(m.FactEvent).where(m.FactEvent.site == "WBK-BAU")) == n_bau
    repo.delete_production(s, ["WBK-BAU"])
    s.commit()
    assert s.scalar(select(func.count()).select_from(m.FactEvent).where(m.FactEvent.site == "WBK-BAU")) == 0
    assert s.scalar(select(func.count()).select_from(m.FactEvent).where(m.FactEvent.site == "WBK-MAS")) > 0
    assert s.get(m.Upload, up.id) is not None                       # other sites keep the upload
    # a month range that does not hold the file's month deletes nothing
    month = up.month
    assert repo.delete_production(s, None, month.replace(year=month.year + 1), None, dry_run=True)["versions"] == 0
    repo.delete_production(s)                                       # the rest: upload record goes too
    s.commit()
    assert s.get(m.Upload, up.id) is None
    assert s.scalar(select(func.count()).select_from(m.FactEvent)) == 0


def test_delete_hourly_and_population(db_session):
    s = db_session
    for site, day in (("WBK-BAU", 1), ("WBK-BAU", 20), ("WBK-MAS", 1)):
        sh = m.HourlyShift(site=site, date=dt.date(2026, 9, day), shift="DS", updated_by="t")
        s.add(sh)
        s.flush()
        s.add(m.HourlyRow(shift_id=sh.id, line=1, loader="WEX015", material="OB - FreeDig", material_group="OB",
                          hauler_model="777E", muatan=41))
    units = pd.DataFrame({"unit_id": ["WHT026"], "type": ["Hauling"], "description": [None], "model": ["777E"],
                          "manufacturer": [None], "site": ["WBK-BAU"]})
    repo.save_population(s, units, dt.date(2026, 1, 1), "a.xlsx", "a" * 64, None)
    repo.save_population(s, units, dt.date(2026, 9, 1), "b.xlsx", "b" * 64, None)
    s.commit()
    got = repo.delete_hourly(s, ["WBK-BAU"], dt.date(2026, 9, 1), dt.date(2026, 9, 10))
    assert got == {"shifts": 1, "lines": 1}
    left = s.execute(select(m.HourlyShift.site, m.HourlyShift.date)).all()
    assert sorted(left) == [("WBK-BAU", dt.date(2026, 9, 20)), ("WBK-MAS", dt.date(2026, 9, 1))]
    assert repo.delete_population(s, dt.date(2026, 6, 1), None) == {"versions": 1, "units": 1}
    assert [v.effective_from for v in s.scalars(select(m.PopulationVersion))] == [dt.date(2026, 1, 1)]


@pytest.mark.parametrize("name", ["Hourly", "Hourly Production"])
def test_hourly_template_old_and_new_sheet_name(name):
    lf = pd.DataFrame({"material": ["OB - FreeDig"], "hauler_model": ["777E"], "load": [41.0]})
    tg = pd.DataFrame({"unit_id": ["WEX015"], "material_group": ["OB"], "target_per_hour": [800.0]})
    data = H.build_template("WBK-BAU", dt.date(2026, 9, 27), "DS", lf, tg)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames[0] == "Hourly Production"
    wb["Hourly Production"].title = name
    buf = io.BytesIO()
    wb.save(buf)
    hf = H.parse_template(buf.getvalue())
    assert (hf.site, hf.shift) == ("WBK-BAU", "DS")


@pytest.mark.parametrize("page", ["pages/home.py", "pages/admin/delete_data.py", "pages/data/upload.py",
                                  "pages/admin/unit_population.py", "pages/data/approval.py",
                                  "pages/admin/targets_plan.py", "pages/admin/hourly_targets.py",
                                  "pages/admin/hourly_setup.py"])
def test_dataset_pages_render_for_admin(db_session, sample_bytes, monkeypatch, page):
    from streamlit.testing.v1 import AppTest

    from auth import security
    from auth.access import ADMIN, load_user
    from core.config import database_url

    from .conftest import APP_DIR

    s = db_session
    ing.ingest(s, sample_bytes, "Eq.Event.xlsb", username="t")
    security.create_user(s, "adm", "Admin", ADMIN, "Tambang-2026x")
    s.commit()
    monkeypatch.setenv("DATABASE_URL", database_url(test=True))
    at = AppTest.from_file(str(APP_DIR / page), default_timeout=180)
    at.session_state["user"] = load_user(s, "adm")
    at.run()
    assert not at.exception, [e.value for e in at.exception]


def test_data_officer_changes_only_units_of_own_sites():
    cols = ["unit_id", "type", "description", "model", "manufacturer", "site"]
    old = pd.DataFrame([("WHT001", "Hauling", None, "777E", None, "WBK-BAU"),
                        ("WHT002", "Hauling", None, "777E", None, "WBK-MAS")], columns=cols)
    new = pd.DataFrame([("WHT001", "Hauling", None, "777F", None, "WBK-BAU"),      # own site: applied
                        ("WHT002", "Hauling", None, "777F", None, "WBK-MAS"),      # other site: ignored
                        ("WHT003", "Hauling", None, "777E", None, "WBK-BAU"),      # added in own site
                        ("WHT004", "Hauling", None, "777E", None, "WBK-MAS")], columns=cols)  # added elsewhere
    keep, ignored = population.restrict(old, new, ["WBK-BAU"])
    got = keep.set_index("unit_id")
    assert sorted(got.index) == ["WHT001", "WHT002", "WHT003"]
    assert got.loc["WHT001", "model"] == "777F" and got.loc["WHT002", "model"] == "777E"
    assert sorted(ignored["unit_id"]) == ["WHT002", "WHT004"]
    # moving a unit out of the user's site touches another site: not applied
    moved = new.assign(site=["WBK-MAS", "WBK-MAS", "WBK-BAU", "WBK-MAS"]).iloc[:1]
    keep, ignored = population.restrict(old.iloc[:1], moved, ["WBK-BAU"])
    assert keep["site"].tolist() == ["WBK-BAU"] and ignored["change"].tolist() == ["moved to WBK-MAS"]


def test_saved_shifts_take_new_targets_on_request(db_session):
    s = db_session
    s.add(m.Site(code="WBK-BAU", name="BAU"))
    sh = m.HourlyShift(site="WBK-BAU", date=dt.date(2026, 9, 1), shift="DS", updated_by="t")
    s.add(sh)
    s.flush()
    s.add(m.HourlyRow(shift_id=sh.id, line=1, loader="WEX015", loader_model="CAT6020B", material="OB - FreeDig",
                      material_group="OB", hauler_model="777E", muatan=41, target_per_hour=800,
                      target_source="default"))
    s.commit()
    s.add(m.HourlyModelTarget(site="WBK-BAU", model="6020B", basis="internal", ob=950))
    s.commit()
    assert repo.recalc_hourly_targets(s, "WBK-BAU", dt.date(2026, 9, 1), dt.date(2026, 9, 30)) == \
        {"shifts": 1, "lines": 1, "changed": 1}
    s.commit()
    row = s.scalar(select(m.HourlyRow))
    assert (row.target_per_hour, row.target_source) == (950, "hourly")
    assert repo.recalc_hourly_targets(s, "WBK-BAU", dt.date(2026, 9, 1), dt.date(2026, 9, 30))["changed"] == 0
