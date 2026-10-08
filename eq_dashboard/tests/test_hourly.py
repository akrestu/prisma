"""Hourly production: production hour, Link Muatan import, resolving rows, template round trip, storage."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
import pytest

from core import dataprod
from core import hourly as H
from core.validate import StructureError
from db import repo

MST = Path(__file__).resolve().parents[2] / "Mst Hourly.xlsx"
LF = pd.DataFrame([("OB - FreeDig", "OB", "777E", 41.0), ("OB - MUD", "OB", "773E", 12.0),
                   ("CG - Coal Getting", "CG", "CWE370Q", 22.5)],
                  columns=["material", "material_group", "hauler_model", "muatan"])
TG = pd.DataFrame([("WEX019", "CAT6020", "OB", 800.0), ("WEX010", "SK200", "CG", 200.0)],
                  columns=["unit_id", "model", "material_group", "target_per_hour"])


def rows(**extra):
    base = {"loader": "WEX019", "operator": "Zainudin", "material": "OB - FreeDig", "hauler_model": "777E",
            "pit": "TKM", "disposal": "IPD Atas", "distance_m": 2600, "r1": 8, "r2": 14, "remark_code": "100",
            "remark": None}
    return pd.DataFrame([{**base, **extra}])


@pytest.mark.parametrize("hhmm,expected", [
    ((6, 0), (dt.date(2026, 9, 26), "DS", 1, "06-07")), ((17, 59), (dt.date(2026, 9, 26), "DS", 12, "17-18")),
    ((18, 0), (dt.date(2026, 9, 26), "NS", 1, "18-19")), ((23, 30), (dt.date(2026, 9, 26), "NS", 6, "23-00")),
])
def test_production_hour(hhmm, expected):
    d, s, k = H.production_hour(dt.datetime(2026, 9, 26, *hhmm))
    assert (d, s, k, H.slot_label(s, k)) == expected


def test_night_shift_after_midnight_belongs_to_previous_date():
    assert H.production_hour(dt.datetime(2026, 9, 27, 2, 30)) == (dt.date(2026, 9, 26), "NS", 9)
    assert H.production_hour(dt.datetime(2026, 9, 27, 5, 59)) == (dt.date(2026, 9, 26), "NS", 12)


@pytest.mark.skipif(not MST.exists(), reason="Mst Hourly.xlsx not found")
def test_link_muatan_from_real_workbook():
    lf, tg = H.parse_link_muatan(MST.read_bytes())
    look = lf.set_index(["material", "hauler_model"])["muatan"]
    assert look[("OB - FreeDig", "777E")] == 41 and look[("CG - Coal Getting", "CWE370Q")] == 22.5
    assert set(lf["material_group"]) == {"OB", "CG"} and (lf["muatan"] > 0).all()
    assert tg.set_index("unit_id").loc["WEX019", "target_per_hour"] == 800 and len(tg) == 22


def test_resolve_fills_load_and_target_and_flags_problems():
    ok = H.resolve(rows(), LF, TG)
    r = ok.rows.iloc[0]
    assert not ok.problems and r["muatan"] == 41 and r["target_per_hour"] == 800 and r["loader_model"] == "CAT6020"
    bad = H.resolve(pd.concat([rows(hauler_model="999X"), rows(r3=75)]), LF, TG)
    assert any("no load per trip" in p for p in bad.problems) and any("between 0 and 20" in p for p in bad.problems)
    empty = H.resolve(rows(loader=None, r1=None, r2=None), LF, TG)       # blank lines are dropped, not errors
    assert empty.rows.empty and not empty.problems


def test_template_round_trip_and_hour_headers():
    tpl = H.build_template("WBK-BAU", dt.date(2026, 9, 26), "NS", LF, TG, H.resolve(rows(), LF, TG).rows, "D. Purba")
    hf = H.parse_template(tpl)
    assert (hf.site, hf.date, hf.shift, hf.coordinator) == ("WBK-BAU", dt.date(2026, 9, 26), "NS", "D. Purba")
    assert hf.rows.loc[0, "loader"] == "WEX019" and pd.isna(hf.rows.loc[0, "r1"])   # lines pre-filled, no trips
    with pytest.raises(StructureError, match="Hourly"):     # a Data_Prod workbook is not an hourly sheet
        H.parse_template(dataprod.build_template(["WBK-BAU"]))


def test_volume_per_hour():
    res = H.resolve(rows(), LF, TG).rows.assign(site="WBK-BAU", date=dt.date(2026, 9, 26), shift="DS")
    long = H.to_long(res)
    by = long.set_index("hour_slot")["volume"]
    assert by["06-07"] == 8 * 41 and by["07-08"] == 14 * 41 and by["08-09"] == 0


def test_save_replace_and_previous_lines(db_session):
    s = db_session
    d = dt.date(2026, 9, 26)
    r = H.resolve(pd.concat([rows(), rows(loader="WEX010", material="CG - Coal Getting", hauler_model="CWE370Q")]),
                  LF, TG).rows
    repo.save_hourly(s, "WBK-BAU", d, "DS", "D. Purba", r, "op1")
    s.commit()
    sh, got = repo.hourly_shift(s, "WBK-BAU", d, "DS")
    assert sh.coordinator == "D. Purba" and len(got) == 2 and got.loc[0, "r2"] == 14
    repo.save_hourly(s, "WBK-BAU", d, "DS", "D. Purba", r.iloc[:1], "op1")      # saving again replaces the sheet
    s.commit()
    assert len(repo.hourly_shift(s, "WBK-BAU", d, "DS")[1]) == 1
    nxt = repo.previous_lines(s, "WBK-BAU", d, "NS")
    assert nxt.loc[0, "loader"] == "WEX019" and pd.isna(nxt.loc[0, "r1"])
    rng = repo.hourly_range(s, ["WBK-BAU"], d, d)
    assert len(rng) == 1 and rng.loc[0, "shift"] == "DS"


UNITS = pd.DataFrame([("WHT026", "Hauling", "DT", "777E", "CAT", "WBK-BAU"),
                      ("WHT027", "Hauling", "DT", "777E", "CAT", "WBK-BAU"),
                      ("WDT017", "Hauling", "DT", "CWE370Q", "UD", "WBK-BAU"),
                      ("WEX019", "Loading", "EX", "CAT6020", "CAT", "WBK-BAU")],
                     columns=["unit_id", "type", "description", "model", "manufacturer", "site"])
OPS = pd.DataFrame([("11001", "Zainudin", "Operator Excavator", True), ("22001", "Budi S.", "Operator Dump Truck", True),
                    ("22002", "Andi P.", "Operator Dump Truck", True)], columns=["nrp", "name", "position", "active"])


def hauler_rows(*lines):
    base = {"loader": "WEX019", "loader_nrp": "11001 - Zainudin", "material": "OB - FreeDig", "pit": "TKM",
            "disposal": "IPD", "distance_m": 2600, "remark_code": None, "remark": None}
    return pd.DataFrame([{**base, **x} for x in lines])


def test_per_hauler_lines_take_model_and_operator_names():
    res = H.resolve(hauler_rows({"hauler": "WHT026", "hauler_nrp": "22001", "r1": 2, "r2": 3},
                                {"hauler": "WHT027", "hauler_nrp": "22002 - Andi P.", "r1": 3}), LF, TG, UNITS, OPS)
    assert not res.problems, res.problems
    r = res.rows.set_index("hauler")
    assert r.loc["WHT026", "hauler_model"] == "777E" and r.loc["WHT026", "muatan"] == 41
    assert r.loc["WHT026", "hauler_operator"] == "Budi S." and r.loc["WHT027", "hauler_nrp"] == "22002"
    assert (r["operator"] == "Zainudin").all() and (r["loader_nrp"] == "11001").all()
    long = H.to_long(res.rows.assign(site="WBK-BAU", date=dt.date(2026, 9, 26), shift="DS"))
    by_op = long.groupby("hauler_operator")["rit"].sum()          # the base of operator KPIs
    assert by_op["Budi S."] == 5 and by_op["Andi P."] == 3


def test_operator_change_mid_shift_and_unknowns():
    res = H.resolve(hauler_rows({"hauler": "WHT026", "hauler_nrp": "22001", "r1": 2, "r2": 2},
                                {"hauler": "WHT026", "hauler_nrp": "22002", "r3": 3},     # new operator, same truck
                                {"hauler": "WHT999", "hauler_nrp": "99999", "r1": 1}), LF, TG, UNITS, OPS)
    assert any("WHT999 is not in the unit population" in p for p in res.problems)
    assert any("99999" in w for w in res.warnings)
    assert not any("Same hauler" in w for w in res.warnings)       # different hours: no overlap warning
    both = H.resolve(hauler_rows({"hauler": "WHT026", "r1": 2}, {"hauler": "WHT026", "r1": 1}), LF, TG, UNITS, OPS)
    assert not both.problems and not any("more than one excavator" in w for w in both.warnings)   # 2 destinations
    assert any("no hauler operator" in w for w in both.warnings)
    moved = H.resolve(hauler_rows({"hauler": "WHT026", "r1": 2}, {"hauler": "WHT026", "loader": "WEX010", "r1": 1}),
                      LF, TG, UNITS, OPS)
    assert any("more than one excavator" in w for w in moved.warnings)
    over = H.resolve(hauler_rows({"hauler": "WHT026", "r1": 12}, {"hauler": "WHT026", "r1": 9}), LF, TG, UNITS, OPS)
    assert any("21 trips in hour 1" in p for p in over.problems)                # 20 per hour over all its lines


def test_per_hauler_template_round_trip():
    lines = H.resolve(hauler_rows({"hauler": "WHT026", "hauler_nrp": "22001"}), LF, TG, UNITS, OPS).rows
    tpl = H.build_template("WBK-BAU", dt.date(2026, 9, 26), "DS", LF, TG, lines, "", UNITS, OPS)
    hf = H.parse_template(tpl)
    r = hf.rows.iloc[0]
    assert (r["loader"], r["hauler"]) == ("WEX019", "WHT026")
    assert H.nrp_of(r["hauler_nrp"]) == "22001" and H.nrp_of(r["loader_nrp"]) == "11001"
    again = H.resolve(hf.rows.assign(r1=4), LF, TG, UNITS, OPS)
    assert not again.problems and again.rows.loc[0, "hauler_operator"] == "Budi S."


def test_template_only_asks_what_the_officer_knows():
    import io as _io

    from openpyxl import load_workbook
    lines = H.resolve(hauler_rows({"hauler": "WHT026", "hauler_nrp": "22001"}), LF, TG, UNITS, OPS).rows
    spare = pd.DataFrame([{"loader": "WEX008", "operator": "Saprin", "material": "OB - FreeDig", "hauler": None,
                           "hauler_model": "777E"}])                      # old per-model line → blank truck row
    tpl = H.build_template("WBK-BAU", dt.date(2026, 9, 28), "NS", LF, TG, pd.concat([lines, spare]), "", UNITS, OPS)
    ws = load_workbook(_io.BytesIO(tpl))["Hourly Production"]
    heads = [c.value for c in ws[H.HEADER_ROW]][:9]
    assert heads == ["Loader", "Operator", "Material", "Hauler ID", "Hauler operator", "PIT", "Disposal",
                     "H distance (m)", "V distance (m)"]
    assert "Hauler model" not in [c.value for c in ws[H.HEADER_ROW]]
    hf = H.parse_template(tpl)
    res = H.resolve(hf.rows.assign(r1=[3, None]), LF, TG, UNITS, OPS)   # the spare row stays unused
    assert not res.problems and len(res.rows) == 1 and res.rows.loc[0, "hauler_model"] == "777E"
    bad = H.resolve(hf.rows.assign(r1=[3, 2]), LF, TG, UNITS, OPS)       # trips but no truck → must be fixed
    assert any("hauler is empty" in p for p in bad.problems)


def test_hauler_model_to_load_class():
    lfm = ["773E", "775E", "777E", "CGE37084R", "CWE370Q"]
    assert H.load_model("777E-KDP", lfm) == "777E" and H.load_model("773E-PRB", lfm) == "773E"
    assert H.load_model("CGE37084R", lfm) == "CGE37084R" and H.load_model("CWE37064R", lfm) is None
    assert H.load_model("CWE37064R", lfm, {"CWE37064R": "CWE370Q"}) == "CWE370Q"
    assert H.suggest_load_model("CWE37064R", lfm) == "CWE370Q" and H.suggest_load_model("775F-DLS", lfm) == "775E"
    assert H.suggest_load_model("AXOR 2528", lfm) is None


def test_dump_trucks_typed_supporting_equipment_count_as_haulers():
    u = pd.DataFrame([("WDT001", "Supporting Equipment", "Hauling 23 Ton"), ("WHT026", "Hauling", "DT"),
                      ("WEX019", "Loading", "EX"), ("WWT001", "Supporting Equipment", "Water Tank")],
                     columns=["unit_id", "type", "description"])
    assert list(u.loc[H.is_hauler(u), "unit_id"]) == ["WDT001", "WHT026"]


def test_mapped_model_and_material_hint():
    units = pd.concat([UNITS, pd.DataFrame([("WDT001", "Supporting Equipment", "Hauling 23 Ton", "CWE37064R", "UD",
                                             "WBK-BAU")], columns=UNITS.columns)])
    rows = hauler_rows({"hauler": "WDT001", "material": "CG - Coal Getting", "r1": 2})
    assert any("has no load per trip" in p for p in H.resolve(rows, LF, TG, units, OPS).problems)
    ok = H.resolve(rows, LF, TG, units, OPS, {"CWE37064R": "CWE370Q"})
    assert not ok.problems and ok.rows.loc[0, "muatan"] == 22.5 and ok.rows.loc[0, "hauler_model"] == "CWE37064R"
    typo = H.resolve(hauler_rows({"hauler": "WHT026", "material": "OB - FreeDigg", "r1": 1}), LF, TG, units, OPS)
    assert any("Did you mean 'OB - FreeDig'" in p for p in typo.problems)


def test_population_falls_back_to_published_data_prod(db_session, sample_bytes):
    from core import ingest as ing
    from db import models as m
    s = db_session
    assert repo.population_for(s, dt.date(2026, 9, 1)) is None
    ing.ingest(s, sample_bytes, "Data_Prod_2026-09.xlsb", username="t")
    for us in s.query(m.UploadSite).filter(m.UploadSite.site_code != "UNMAPPED"):
        ing.publish(s, us, None, "t")
    s.commit()
    units = repo.population_for(s, dt.date(2026, 9, 27))
    assert units is not None and "latest published Production Data" in units.attrs["source"]
    assert units.set_index("unit_id").loc["WHT018", "model"] == "777E-KDP"


def test_load_factors_follow_the_population_models():
    filled, rep = H.expand_to_population(LF, ["777E-KDP", "777E", "CWE37064R", "AXOR 2528"])
    how = rep.set_index("model")["how"]
    assert how["777E-KDP"] == "same model" and how["777E"] == "own values"
    assert how["CWE37064R"].startswith("closest name") and how["AXOR 2528"].startswith("no match")
    look = filled.set_index(["material", "hauler_model"])["muatan"]
    assert look[("OB - FreeDig", "777E-KDP")] == 41 and look[("CG - Coal Getting", "CWE37064R")] == 22.5
    units = pd.DataFrame([("WHT018", "Hauling", "DT", "777E-KDP", "CAT", "WBK-BAU")], columns=UNITS.columns)
    res = H.resolve(hauler_rows({"hauler": "WHT018", "r1": 2}), filled, TG, units, OPS)
    assert not res.problems and res.rows.loc[0, "muatan"] == 41          # found by its own model, no mapping


def test_clean_remarks_per_hour():
    raw = pd.DataFrame([{"hour": "09-10", "loader": "wex019", "hauler": None, "code": "302 - Rain", "remark": None},
                        {"hour": 4, "loader": "WEX019", "hauler": "WHT026", "code": None, "remark": "ban bocor"},
                        {"hour": None, "loader": None, "hauler": None, "code": None, "remark": None}])
    out, problems = H.clean_remarks(raw, "DS", ["WEX019"])
    assert not problems
    assert out[["slot", "loader", "hauler", "code"]].astype(object).where(out.notna(), None).values.tolist() == [
        [4, "WEX019", None, "302"], [4, "WEX019", "WHT026", None]]
    _, problems = H.clean_remarks(pd.DataFrame([{"hour": "19-20", "loader": "WEX999", "code": None}]), "DS",
                                    ["WEX019"])
    assert len(problems) == 3          # not a DS hour, loader not in the shift, no code or text
    assert H.remark_label("302") == "302 - Rain" and H.remark_label("999") == "999"


def test_template_remarks_sheet_round_trip():
    lines = H.resolve(hauler_rows({"hauler": "WHT026", "hauler_nrp": "22001"}), LF, TG, UNITS, OPS).rows
    rm = pd.DataFrame([{"slot": 3, "loader": "WEX019", "hauler": None, "code": "302", "remark": "hujan deras"}])
    tpl = H.build_template("WBK-BAU", dt.date(2026, 9, 26), "DS", LF, TG, lines, "", UNITS, OPS, rm)
    hf = H.parse_template(tpl)
    assert "Remark" not in " ".join(map(str, hf.rows.columns.drop(["remark_code", "remark"])))
    out, problems = H.clean_remarks(hf.remarks, "DS", ["WEX019"])
    assert not problems and out.loc[0, "slot"] == 3 and out.loc[0, "code"] == "302"
    assert out.loc[0, "remark"] == "hujan deras"


def test_save_hourly_keeps_remarks_per_hour(db_session):
    s = db_session
    d = dt.date(2026, 9, 26)
    r = H.resolve(rows(), LF, TG).rows
    rm = pd.DataFrame([{"slot": 2, "loader": "WEX019", "hauler": None, "code": "401", "remark": None}])
    repo.save_hourly(s, "WBK-BAU", d, "DS", "", r, "op1", remarks=rm)
    s.commit()
    got = repo.hourly_remarks(s, "WBK-BAU", d, "DS")
    assert got[["slot", "loader", "code"]].values.tolist() == [[2, "WEX019", "401"]]
    repo.save_hourly(s, "WBK-BAU", d, "DS", "", r, "op1")              # no remarks given: left as they are
    s.commit()
    assert len(repo.hourly_remarks(s, "WBK-BAU", d, "DS")) == 1
    repo.save_hourly(s, "WBK-BAU", d, "DS", "", r, "op1", remarks=rm.iloc[:0])
    s.commit()
    assert repo.hourly_remarks(s, "WBK-BAU", d, "DS").empty


def test_group_remarks_merges_consecutive_hours():
    d = pd.DataFrame({"id": [1, 2, 3, 4, 5], "slot": [3, 4, 5, 7, 4], "loader": ["A", "A", "A", "A", "B"],
                      "hauler": [None] * 5, "code": ["302"] * 4 + ["401"], "remark": [None] * 5})
    g = H.group_remarks(d)
    assert g[["slot_from", "slot_to", "loader", "code"]].values.tolist() == [[3, 5, "A", "302"], [4, 4, "B", "401"],
                                                                            [7, 7, "A", "302"]]
    assert g.loc[0, "ids"] == [1, 2, 3]
    assert H.span_label("DS", 3, 5) == "08-09 – 10-11" and H.span_label("NS", 1, 1) == "18-19"
    assert H.group_remarks(d.iloc[:0]).empty


def test_add_and_delete_remark_without_saving_the_grid(db_session):
    s = db_session
    d = dt.date(2026, 9, 27)
    n = repo.add_hourly_remark(s, "WBK-BAU", d, "NS", 4, 2, "WEX019", "401", " hose bocor ", "sb1")
    s.commit()
    assert n == 3                                               # 19-20 … 21-22, given in either order
    got = repo.hourly_remarks(s, "WBK-BAU", d, "NS")
    assert got["slot"].tolist() == [2, 3, 4] and set(got["remark"]) == {"hose bocor"}
    sh, rows = repo.hourly_shift(s, "WBK-BAU", d, "NS")
    assert sh is not None and rows.empty                        # the shift exists, no trip lines needed
    ev = H.group_remarks(got)
    repo.delete_hourly_remarks(s, ev.loc[0, "ids"])
    s.commit()
    assert repo.hourly_remarks(s, "WBK-BAU", d, "NS").empty


def test_remark_tags_are_words_never_numbers():
    assert [H.remark_tag(c) for c in ("302", "401", "455", "599", None)] == ["RAIN", "BD-L", "BD", "MAINT", "NOTE"]
    assert all(not t.replace("-", "").isdigit() for t in H.REMARK_TAGS.values())
    assert [H.remark_category(c) for c in ("401", "302", "501", "100", None)] == ["down", "delay", "maint", "info",
                                                                                   "note"]
