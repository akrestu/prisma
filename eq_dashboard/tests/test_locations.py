"""PIT and disposal master: checking the list, the shift lines against it, the template drop-downs, saving."""
from __future__ import annotations

import datetime as dt
import io

import pandas as pd
from openpyxl import Workbook, load_workbook

from core import hourly as H
from core import locations as LC
from db import repo
from tests.test_hourly import LF, OPS, TG, UNITS, hauler_rows

D = dt.date(2026, 10, 7)
LOCS = pd.DataFrame([("PIT", "TKM", "OB", True), ("PIT", "Pit 3", "CG", True), ("PIT", "Old Pit", "OB", False),
                     ("DISPOSAL", "IPD Atas", "OB", True), ("DISPOSAL", "ROM A", "CG", True),
                     ("DISPOSAL", "ROM B", "CG", True), ("DISPOSAL", "Old Dump", "OB", False)], columns=LC.COLS)


def test_checks_on_the_master():
    d, p = LC.clean(pd.DataFrame([("PIT", "Pit 3", "cg", None), ("pit", "pit 3", "CG", "Yes"),
                                  ("DISPOSAL", "Pit 3", "CG", True),          # same name, other type: fine
                                  ("DUMP", "X", "OB", True), ("PIT", "Y", "MUD", "No")], columns=LC.COLS))
    assert any("Listed twice" in x and "Pit 3" in x for x in p)
    assert any("X: type must be PIT or Disposal" in x for x in p) and any("Y: material must be OB or CG" in x for x in p)
    assert d.loc[0, "active"] and not d.loc[4, "active"] and d.loc[1, "kind"] == "PIT"
    assert LC.names(LOCS, "PIT") == ["Pit 3", "TKM"] and "Old Dump" not in LC.names(LOCS, "DISPOSAL")


def test_lines_keep_typed_pit_disposal_and_distances():
    lines = hauler_rows({"hauler": "WHT026", "disposal": "ipd atas", "pit": "tkm", "distance_m": 1, "r1": 2},
                        {"hauler": "WHT027", "disposal": "IPD Atas", "loader": "WEX010", "r1": 1})
    res = H.resolve(lines, LF, TG, UNITS, OPS, locations=LOCS)
    assert not res.problems, res.problems
    r = res.rows.set_index("hauler")
    assert tuple(r.loc["WHT026", ["pit", "disposal", "distance_m"]]) == ("TKM", "IPD Atas", 1)   # spelled as listed
    assert pd.isna(r.loc["WHT026", "dist_v"]) and any("no H or V distance" in w for w in res.warnings)


def test_pit_and_disposal_required_active_and_of_the_right_material():
    res = H.resolve(hauler_rows({"hauler": "WHT026", "disposal": None, "pit": None, "r1": 2},
                                {"hauler": "WHT027", "disposal": "Old Dump", "pit": "Old Pit", "r1": 1},
                                {"hauler": "WHT026", "disposal": "ROM A", "pit": "Pit 3", "r2": 1},
                                {"hauler": "WHT027", "disposal": None, "pit": None}),   # no trips: nothing needed
                    LF, TG, UNITS, OPS, locations=LOCS)
    p = res.problems
    assert sum("pick the disposal" in x for x in p) == 1 and sum("pick the PIT" in x for x in p) == 1
    assert any("Old Dump is not an active disposal" in x for x in p)
    assert any("Old Pit is not an active PIT" in x for x in p)
    assert any("ROM A is a CG disposal but the material is OB - FreeDig" in x for x in p)
    assert any("Pit 3 is a CG PIT but the material is OB - FreeDig" in x for x in p)


def test_hauler_to_two_disposals_in_one_hour():
    cg = {"hauler": "WDT017", "material": "CG - Coal Getting", "loader": "WEX010", "pit": "Pit 3"}
    res = H.resolve(hauler_rows({**cg, "disposal": "ROM A", "r1": 2}, {**cg, "disposal": "ROM B", "r1": 1}),
                    LF, TG, UNITS, OPS, locations=LOCS)
    assert not res.problems, res.problems
    long = H.to_long(res.rows.assign(site="WBK-BAU", date=D, shift="DS"))
    assert long.loc[long["slot"] == 1, "volume"].sum() == 3 * 22.5          # trips × load, both disposals


def _old_template(heads: list[str], values: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = H.SHEET
    for r, (k, v) in enumerate([("Site", "WBK-BAU"), ("Date", D), ("Shift", "DS"), ("Shift boss", "")], start=2):
        ws.cell(r, 1, k)
        ws.cell(r, 2, v)
    for j, h in enumerate(heads + H.SLOTS["DS"], start=1):
        ws.cell(H.HEADER_ROW, j, h)
    for j, v in enumerate(values, start=1):
        ws.cell(H.HEADER_ROW + 1, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_template_lists_pit_and_disposal_and_older_templates_still_read():
    lines = H.resolve(hauler_rows({"hauler": "WHT026", "disposal": "IPD Atas", "pit": "TKM"}), LF, TG, UNITS, OPS,
                      locations=LOCS).rows.assign(distance_m=2450.0, dist_v=35.0)
    tpl = H.build_template("WBK-BAU", D, "DS", LF, TG, lines, "", UNITS, OPS, locations=LOCS)
    wb = load_workbook(io.BytesIO(tpl))
    assert [c.value for c in wb[H.SHEET][H.HEADER_ROW]][:9] == H.INPUT_COLS
    assert [c.value for c in wb["Lists"]["H"] if c.value] == ["Pit 3", "TKM"]          # active PITs only
    assert "Old Dump" not in [c.value for c in wb["Lists"]["G"]]
    hf = H.parse_template(tpl)
    assert tuple(hf.rows.loc[0, ["pit", "disposal", "distance_m", "dist_v"]]) == ("TKM", "IPD Atas", 2450, 35)
    v3 = H.parse_template(_old_template(H.INPUT_COLS_V3, ["WEX019", None, "OB - FreeDig", "WHT026", None,
                                                          "IPD Atas", 2600, 40, 3])).rows.iloc[0]
    assert (v3["disposal"], v3["distance_m"], v3["dist_v"], v3["r1"]) == ("IPD Atas", 2600, 40, 3)
    v1 = H.parse_template(_old_template(H.INPUT_COLS_V1, ["WEX019", None, "OB - FreeDig", "WHT026", None, "TKM",
                                                          "IPD", 2600, 3])).rows.iloc[0]
    assert (v1["pit"], v1["disposal"], v1["distance_m"], v1["r1"]) == ("TKM", "IPD", 2600, 3)


def test_save_master_and_shift(db_session):
    s = db_session
    assert repo.save_haul_locations(s, "WBK-BAU", LOCS) == len(LOCS)
    s.commit()
    got = repo.haul_locations(s, "WBK-BAU")
    assert len(got) == len(LOCS) and set(got["kind"]) == {"PIT", "DISPOSAL"}
    rows = H.resolve(hauler_rows({"hauler": "WHT026", "disposal": "IPD Atas", "pit": "TKM", "dist_v": 40, "r1": 2}),
                     LF, TG, UNITS, OPS, locations=got).rows
    repo.save_hourly(s, "WBK-BAU", D, "DS", "", rows, "op1")
    repo.save_haul_locations(s, "WBK-BAU", LOCS[LOCS["name"] != "TKM"])     # removed later: the shift keeps it
    s.commit()
    saved = repo.hourly_shift(s, "WBK-BAU", D, "DS")[1].iloc[0]
    assert (saved["pit"], saved["disposal"], saved["dist_v"]) == ("TKM", "IPD Atas", 40)
