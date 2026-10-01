"""Load Factors workbook (replaces Mst Hourly): template round trip, matrix ↔ long, refusals."""
from __future__ import annotations

import io

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from core import load_factors as LF
from core.validate import StructureError

LOAD = pd.DataFrame({"material": ["OB - FreeDig", "OB - FreeDig", "CG - Coal Getting"],
                     "material_group": ["OB", "OB", "CG"], "hauler_model": ["777E", "CWE370", "CWE370"],
                     "muatan": [41.0, 24.0, 22.5]})


def _key(df):
    return df.sort_values(["material", "hauler_model"])[["material", "material_group", "hauler_model", "muatan"]] \
        .reset_index(drop=True)


def test_matrix_round_trip_with_population_columns_first():
    g = LF.wide(LOAD, ["CWE370", "777E", "HD785"])
    assert list(g.columns) == ["material", "CWE370", "777E", "HD785"]       # an empty model column is kept
    pd.testing.assert_frame_equal(_key(LF.long(g)), _key(LOAD))


def test_template_round_trip():
    data = LF.build_template("WBK-BAU", LOAD, ["777E", "CWE370"])
    assert load_workbook(io.BytesIO(data)).sheetnames[0] == "Load Factors"
    got = LF.parse(data)
    assert got.site == "WBK-BAU" and got.source == "Load Factors" and got.legacy_targets == 0
    pd.testing.assert_frame_equal(_key(got.load), _key(LOAD))
    assert LF.file_name("WBK-BAU") == "Load_Factors_WBK-BAU.xlsx"


def test_material_must_be_ob_or_cg():
    data = LF.build_template("WBK-BAU", LOAD, ["777E"])
    wb = load_workbook(io.BytesIO(data))
    wb["Load Factors"].cell(LF.HEADER_ROW + 1, 1, "Sand")
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(StructureError, match="OB or CG"):
        LF.parse(buf.getvalue())


def test_unknown_workbook_is_refused():
    wb = Workbook()
    wb.active.title = "Something"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(StructureError, match="Load Factors"):
        LF.parse(buf.getvalue())
