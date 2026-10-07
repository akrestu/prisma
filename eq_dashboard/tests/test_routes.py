"""Destinations (Tujuan) and routes: checking, the route in force on a date, the workbook, and the shift lines."""
from __future__ import annotations

import datetime as dt
import io

import pandas as pd
from openpyxl import Workbook

from core import hourly as H
from core import routes as RT
from db import repo
from tests.test_hourly import LF, OPS, TG, UNITS, hauler_rows

D = dt.date(2026, 10, 7)
DEST = pd.DataFrame([("IPD Atas", "OB", True), ("ROM A", "CG", True), ("ROM B", "CG", True),
                     ("Old Dump", "OB", False)], columns=RT.DEST_COLS)
ROUTES = pd.DataFrame([("WEX019", "IPD Atas", "TKM", 2600.0, 40.0, dt.date(2026, 10, 1)),
                       ("WEX019", "IPD Atas", "TKM", 2900.0, 45.0, dt.date(2026, 10, 10)),   # planned change
                       ("WEX010", "ROM A", "Pit 3", 1200.0, 35.0, dt.date(2026, 9, 1))],
                      columns=RT.ROUTE_COLS)


def test_route_in_force_on_a_date():
    now = RT.routes_at(ROUTES, D).set_index(["loader", "destination"])
    assert now.loc[("WEX019", "IPD Atas"), "dist_h"] == 2600 and now.loc[("WEX010", "ROM A"), "pit"] == "Pit 3"
    assert RT.routes_at(ROUTES, dt.date(2026, 10, 12)).set_index("loader").loc["WEX019", "dist_h"] == 2900
    assert len(RT.current(ROUTES, D)) == 3                              # in force + planned


def test_saving_the_grid_keeps_replaced_routes_as_history():
    old = pd.concat([ROUTES, pd.DataFrame([("WEX019", "IPD Atas", "TKM", 2000.0, 30.0, dt.date(2026, 9, 1))],
                                          columns=RT.ROUTE_COLS)], ignore_index=True)
    edited = RT.current(old, D).assign(dist_v=50.0)                     # the grid shows only in force + planned
    out = RT.merge_routes(old, edited, D)
    assert len(out) == 4
    assert sorted(out.loc[out["valid_from"] == dt.date(2026, 9, 1), "dist_h"]) == [1200.0, 2000.0]
    assert out.loc[out["valid_from"] == dt.date(2026, 10, 1), "dist_v"].iloc[0] == 50


def test_checks_on_destinations_and_routes():
    d, p = RT.clean_destinations(pd.DataFrame([("ROM A", "cg", None), ("rom a", "CG", "Yes"), ("X", "MUD", "No")],
                                              columns=RT.DEST_COLS))
    assert d.loc[0, "material_group"] == "CG" and d.loc[0, "active"] and not d.loc[2, "active"]
    assert any("listed twice" in x for x in p) and any("X: material must be OB or CG" in x for x in p)
    r, p = RT.clean_routes(pd.DataFrame([("wex 019", "ipd atas", None, 2600, None, None),
                                         ("WEX019", "Nowhere", None, 99_999, None, None)], columns=RT.ROUTE_COLS),
                           DEST, D)
    assert (r.loc[0, "loader"], r.loc[0, "destination"], r.loc[0, "valid_from"]) == ("WEX019", "IPD Atas", D)
    assert any("Nowhere is not in the destinations" in x for x in p) and any("50,000" in x for x in p)


def test_workbook_round_trip():
    tpl = RT.build_template("WBK-BAU", DEST, ROUTES, ["WEX010", "WEX019"])
    f = RT.parse(tpl, D)
    assert f.site == "WBK-BAU" and not f.problems
    assert list(f.destinations["name"]) == list(DEST["name"]) and not f.destinations.loc[3, "active"]
    assert len(f.routes) == 3 and f.routes.loc[1, "valid_from"] == dt.date(2026, 10, 10)


def test_shift_lines_take_pit_and_distances_from_the_route():
    lines = hauler_rows({"hauler": "WHT026", "disposal": "ipd atas", "pit": "typed", "distance_m": 1, "r1": 2},
                        {"hauler": "WHT027", "disposal": "IPD Atas", "loader": "WEX010", "r1": 1})
    res = H.resolve(lines, LF, TG, UNITS, OPS, destinations=DEST, routes=RT.routes_at(ROUTES, D))
    assert not res.problems, res.problems
    r = res.rows.set_index("hauler")
    assert (r.loc["WHT026", "disposal"], r.loc["WHT026", "pit"], r.loc["WHT026", "distance_m"],
            r.loc["WHT026", "dist_v"]) == ("IPD Atas", "TKM", 2600, 40)
    assert pd.isna(r.loc["WHT027", "distance_m"]) and pd.isna(r.loc["WHT027", "pit"])    # no route: no stale copy
    assert any("No route yet for: WEX010 → IPD Atas" in w for w in res.warnings)


def test_destination_is_required_active_and_of_the_right_material():
    res = H.resolve(hauler_rows({"hauler": "WHT026", "disposal": None, "r1": 2},
                                {"hauler": "WHT027", "disposal": "Old Dump", "r1": 1},
                                {"hauler": "WHT026", "disposal": "ROM A", "r2": 1},
                                {"hauler": "WHT027", "disposal": None}),          # no trips: no destination needed
                    LF, TG, UNITS, OPS, destinations=DEST, routes=RT.routes_at(ROUTES, D))
    assert sum("pick the destination" in p for p in res.problems) == 1
    assert any("Old Dump is not an active destination" in p for p in res.problems)
    assert any("ROM A is a CG destination but the material is OB - FreeDig" in p for p in res.problems)


def test_hauler_to_two_destinations_in_one_hour():
    two = hauler_rows({"hauler": "WDT017", "material": "CG - Coal Getting", "loader": "WEX010", "disposal": "ROM A",
                       "r1": 2}, {"hauler": "WDT017", "material": "CG - Coal Getting", "loader": "WEX010",
                                  "disposal": "ROM B", "r1": 1})
    res = H.resolve(two, LF, TG, UNITS, OPS, destinations=DEST, routes=RT.routes_at(ROUTES, D))
    assert not res.problems, res.problems
    long = H.to_long(res.rows.assign(site="WBK-BAU", date=D, shift="DS"))
    assert long.loc[long["slot"] == 1, "volume"].sum() == 3 * 22.5          # trips × load, both destinations


def test_new_template_asks_destination_and_old_template_still_reads():
    lines = H.resolve(hauler_rows({"hauler": "WHT026", "disposal": "IPD Atas"}), LF, TG, UNITS, OPS,
                      destinations=DEST, routes=RT.routes_at(ROUTES, D)).rows
    hf = H.parse_template(H.build_template("WBK-BAU", D, "DS", LF, TG, lines, "", UNITS, OPS, destinations=DEST))
    assert hf.rows.loc[0, "disposal"] == "IPD Atas" and "distance_m" not in hf.rows
    wb = Workbook()                                                      # a template made before destinations
    ws = wb.active
    ws.title = H.SHEET
    for r, (k, v) in enumerate([("Site", "WBK-BAU"), ("Date", D), ("Shift", "DS"), ("Shift boss", "")], start=2):
        ws.cell(r, 1, k)
        ws.cell(r, 2, v)
    for j, h in enumerate(H.INPUT_COLS_V1 + H.SLOTS["DS"], start=1):
        ws.cell(H.HEADER_ROW, j, h)
    for j, v in enumerate(["WEX019", None, "OB - FreeDig", "WHT026", None, "TKM", "IPD", 2600, 3], start=1):
        ws.cell(H.HEADER_ROW + 1, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    old = H.parse_template(buf.getvalue()).rows.iloc[0]
    assert (old["pit"], old["disposal"], old["distance_m"], old["r1"]) == ("TKM", "IPD", 2600, 3)


def test_save_keeps_a_used_destination_as_inactive(db_session):
    s = db_session
    repo.save_haul_setup(s, "WBK-BAU", DEST, ROUTES)
    s.commit()
    repo.save_haul_setup(s, "WBK-BAU", DEST[DEST["name"] != "ROM A"])     # ROM A is still used by a route
    s.commit()
    got = repo.haul_destinations(s, "WBK-BAU").set_index("name")
    assert not got.loc["ROM A", "active"] and len(repo.haul_routes(s, "WBK-BAU")) == 3
    rows = H.resolve(hauler_rows({"hauler": "WHT026", "disposal": "IPD Atas", "r1": 2}), LF, TG, UNITS, OPS,
                     destinations=got.reset_index(), routes=RT.routes_at(repo.haul_routes(s, "WBK-BAU"), D)).rows
    repo.save_hourly(s, "WBK-BAU", D, "DS", "", rows, "op1")
    s.commit()
    assert repo.hourly_shift(s, "WBK-BAU", D, "DS")[1].loc[0, "dist_v"] == 40
